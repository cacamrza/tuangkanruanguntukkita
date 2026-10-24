import hashlib
import io
import re
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from docx import Document

st.set_page_config(page_title="Ruang Anggaran — Dashboard RAB & TOR",page_icon="balmon.png", layout="wide")
ROOT = Path(__file__).resolve().parent
COLORS = ["#092759", "#087cb5", "#3d91b8", "#7a9bb5", "#a8bbc9", "#c3ccd4"]
MAX_UPLOAD_SIZE = 20 * 1024 * 1024
BUDGET_HEADERS = [
    "_Jenis",
    "Kode Anggaran",
    "Uraian",
    "Pagu Penuh",
    "Pagu Efektif",
    "Pagu Terblokir",
    "Realisasi",
    "Sisa Anggaran",
    "Status Pasangan",
    "Induk",
]
BUDGET_LEVEL_LABELS = {
    "PROGRAM": "Program",
    "KEGIATAN": "Kegiatan",
    "KRO": "KRO",
    "RO": "RO",
    "KOMPONEN": "Komponen",
    "SUBKOMPONEN": "Subkomponen",
    "DETIL": "Detil belanja",
    "ITEM": "Rincian item",
}


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def to_number(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)):
        return float(value) if pd.notna(value) else None
    text = re.sub(r"[^\d,.-]", "", str(value)).strip()
    if not text or re.fullmatch(r"-+", text):
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        if re.fullmatch(r"-?\d{1,3}(?:,\d{3})+", text):
            text = text.replace(",", "")
        else:
            head, tail = text.rsplit(",", 1)
            text = head + "." + tail
    elif "." in text:
        if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", text):
            text = text.replace(".", "")
        else:
            head, tail = text.rsplit(".", 1)
            text = head + tail if len(tail) == 3 and head.lstrip("-").isdigit() else text
    try:
        value = float(text)
        return value if pd.notna(value) else None
    except ValueError:
        return None


def trim_rows(rows):
    result = [list(row) for row in rows]
    while result and all(value is None or clean_text(value) == "" for value in result[-1]):
        result.pop()
    return result


def split_text_rows(text):
    rows = []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if not line.strip():
            continue
        delimiter = "\t" if "\t" in line else re.compile(r" {2,}")
        rows.append([part.strip() or None for part in delimiter.split(line.strip())])
    return trim_rows(rows)


def parse_rtf_rows(content):
    text = content.decode("cp1252", errors="replace")
    text = re.sub(
        r"\\'([0-9a-fA-F]{2})",
        lambda match: bytes.fromhex(match.group(1)).decode("cp1252"),
        text,
    )
    text = re.sub(r"\\u(-?\d+)\??", lambda match: chr(int(match.group(1)) % 65536), text)
    text = re.sub(r"\\(?:par[d]?|row|line)\b ?", "\n", text)
    text = re.sub(r"\\(?:cell|tab)\b ?", "\t", text)
    text = re.sub(r"\\~", " ", text)
    text = re.sub(r"\\([\\{}])", r"\1", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = text.replace("{", "").replace("}", "")
    return split_text_rows(text)


def parse_sp2d_rows(rows):
    joined = " ".join(clean_text(value) for row in rows for value in row)
    is_sp2d = (
        re.search(r"LAPORAN REALISASI SP2D", joined, re.I)
        and any(re.fullmatch(r"Lock Pagu", clean_text(value), re.I) for row in rows for value in row)
        and any(re.fullmatch(r"Uraian", clean_text(value), re.I) for row in rows for value in row)
    )
    if not is_sp2d:
        return None
    normalized = [["Kode Anggaran", "Uraian", "Pagu", "Sisa Anggaran"]]
    for row in rows:
        padded = list(row) + [None] * max(0, 9 - len(row))
        code = clean_text(padded[0])
        description = clean_text(padded[2])
        amount = to_number(padded[8])
        remaining = to_number(padded[6])
        if re.fullmatch(r"(?:\d{3}|\d{6})", code) and description and amount and amount > 0:
            if remaining is not None and remaining >= 0:
                normalized.append([code, description, amount, remaining])
    if len(normalized) == 1:
        raise ValueError("Kolom pagu dan sisa pada laporan SP2D tidak dapat dipastikan.")
    return normalized


def parse_sp2d_text(pages_text):
    if not any(re.search(r"LAPORAN REALISASI SP2D", text, re.I) for text in pages_text):
        return None

    code_pattern = re.compile(r"^\s*(?P<code>\d{3}|\d{6})\.?\s+(?P<description>.*)$")
    amount_pattern = re.compile(r"(?<![\w,])\d{1,3}(?:,\d{3})+(?![\w,])")
    number_pattern = re.compile(r"(?<![\w,])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\w,])")
    normalized = [["Kode Anggaran", "Uraian", "Pagu", "Sisa Anggaran"]]

    for page_text in pages_text:
        pending = None
        wrapped_row = None
        for line in page_text.splitlines():
            text = line.strip()
            if not text:
                continue
            if re.search(r"^(?:Uraian\s+Pagu|Realisasi TA|Periode Lalu|SISA\s+ANGGARAN)", text, re.I):
                pending = None
                wrapped_row = None
                continue

            code_match = code_pattern.match(line)
            amount_match = amount_pattern.search(line)
            if code_match:
                wrapped_row = None
                code = code_match.group("code")
                description_start = code_match.start("description")
                if amount_match:
                    amounts = number_pattern.findall(line[amount_match.start():])
                    if len(amounts) >= 6:
                        normalized.append([
                            code,
                            line[description_start:amount_match.start()].strip(),
                            to_number(amounts[0]),
                            to_number(amounts[-1]),
                        ])
                    pending = None
                else:
                    pending = [code, code_match.group("description").strip()]
                continue

            if amount_match:
                amounts = number_pattern.findall(line[amount_match.start():])
                if pending and len(amounts) >= 6:
                    normalized.append([
                        pending[0],
                        pending[1],
                        to_number(amounts[0]),
                        to_number(amounts[-1]),
                    ])
                    wrapped_row = normalized[-1]
                    pending = None
                else:
                    pending = None
                    wrapped_row = None
                continue

            if pending:
                pending[1] = "{} {}".format(pending[1], text).strip()
            elif wrapped_row:
                wrapped_row[1] = "{} {}".format(wrapped_row[1], text).strip()

    valid_rows = [
        row for row in normalized[1:]
        if row[1] and row[2] is not None and row[2] > 0 and row[3] is not None and row[3] >= 0
    ]
    if not valid_rows:
        raise ValueError("Rincian SP2D pada PDF terdeteksi, tetapi nilai pagu dan sisa tidak terbaca.")
    return [normalized[0]] + valid_rows


def parse_rincian_kertas_text(pages_text):
    if not any(re.search(r"RINCIAN KERTAS KERJA SATKER", text, re.I) for text in pages_text):
        return None

    amount_pattern = re.compile(r"(?<![\w,])\d{1,3}(?:,\d{3})+(?![\w,])")
    patterns = (
        ("PROGRAM", re.compile(r"^(?P<code>\d{3}\.\d{2}\.[A-Z0-9]{2})\s+(?P<description>.*)$", re.I)),
        ("RO", re.compile(r"^(?P<code>\d{4}\.[A-Z0-9]{3}\.[A-Z0-9]{3})\s+(?P<description>.*)$", re.I)),
        ("KRO", re.compile(r"^(?P<code>\d{4}\.[A-Z0-9]{3})\s+(?P<description>.*)$", re.I)),
        ("KEGIATAN", re.compile(r"^(?P<code>\d{4})\s+(?P<description>.*)$", re.I)),
        ("DETIL", re.compile(r"^(?P<code>\d{6})\s+(?P<description>.*)$")),
        ("KOMPONEN", re.compile(r"^(?P<code>\d{3})\s+(?P<description>.*)$")),
        ("SUBKOMPONEN", re.compile(r"^(?P<code>[A-Z])\s+(?P<description>.*)$")),
    )
    parent_levels = {
        "PROGRAM": None,
        "KEGIATAN": "PROGRAM",
        "KRO": "KEGIATAN",
        "RO": "KRO",
        "KOMPONEN": "RO",
        "SUBKOMPONEN": "KOMPONEN",
        "DETIL": "SUBKOMPONEN",
    }
    descendants = {
        "PROGRAM": ("KEGIATAN", "KRO", "RO", "KOMPONEN", "SUBKOMPONEN", "DETIL"),
        "KEGIATAN": ("KRO", "RO", "KOMPONEN", "SUBKOMPONEN", "DETIL"),
        "KRO": ("RO", "KOMPONEN", "SUBKOMPONEN", "DETIL"),
        "RO": ("KOMPONEN", "SUBKOMPONEN", "DETIL"),
        "KOMPONEN": ("SUBKOMPONEN", "DETIL"),
        "SUBKOMPONEN": ("DETIL",),
        "DETIL": (),
    }
    total_allocation = None
    blocked_total = 0
    hierarchy_rows = []
    current = {}

    for page_text in pages_text:
        for line in page_text.splitlines():
            if total_allocation is None and re.search(r"\bALOKASI\b", line, re.I):
                allocation_amounts = list(amount_pattern.finditer(line))
                if allocation_amounts:
                    total_allocation = to_number(allocation_amounts[-1].group(0))

            text = line.strip()
            if not text or re.match(
                r"^(?:RINCIAN KERTAS|KEMEN/LEMB|UNIT ORG|UNIT KERJA|ALOKASI|Halaman\s*:|PROGRAM/|KODE\s*$|SUBKOMP/|DETIL VOLUME|\([1-6]\)|Lokasi\s*:|\(KPPN\.)",
                text,
                re.I,
            ):
                continue

            if re.match(r"^[-•]\s+", text):
                amounts = list(amount_pattern.finditer(text))
                if text.endswith("*") and amounts:
                    blocked = to_number(amounts[-1].group(0))
                    blocked_total += blocked or 0
                    for ancestor in current.values():
                        ancestor[5] = (ancestor[5] or 0) + (blocked or 0)
                continue

            matched = next(
                ((kind, code_match) for kind, pattern in patterns if (code_match := pattern.match(text))),
                None,
            )
            if matched:
                kind, code_match = matched
                code = clean_text(code_match.groupdict().get("code")).upper()
                content = code_match.group("description")
                amounts = list(amount_pattern.finditer(content))
                description = content[:amounts[0].start()].strip() if amounts else content.strip()
                description = re.sub(r"\s+\d+(?:[.,]\d+)?\s+\S+\s*$", "", description).strip()
                amount = to_number(amounts[-1].group(0)) if amounts else None
                parent_level = parent_levels[kind]
                parent = current.get(parent_level) if parent_level else None
                row = [
                    kind,
                    code,
                    description,
                    amount,
                    None,
                    None,
                    None,
                    None,
                    "Menunggu Laporan FA" if kind == "RO" else "Rincian Kertas Kerja",
                    parent[1] if parent else None,
                ]
                hierarchy_rows.append(row)
                current[kind] = row
                for descendant in descendants[kind]:
                    current.pop(descendant, None)
                continue

            if current and not amount_pattern.search(text) and not re.match(r"^[A-Z]\s", text):
                last_row = hierarchy_rows[-1]
                if not text.startswith(("(", "-", "•")):
                    last_row[2] = "{} {}".format(last_row[2], text).strip()

    rows = [BUDGET_HEADERS]
    rows.append([
        "TOTAL", "TOTAL", "ALOKASI", total_allocation, None, blocked_total, None, None, "TOTAL", None,
    ])
    rows.extend(hierarchy_rows)
    if total_allocation is None and not any(row[0] == "RO" for row in hierarchy_rows):
        raise ValueError("PDF Rincian Kertas Kerja terdeteksi, tetapi nilai alokasi dan kode RO tidak terbaca.")
    return rows


def parse_fa_text(pages_text):
    if not any(re.search(r"LAPORAN REALISASI SP2D", text, re.I) for text in pages_text):
        return None

    number_pattern = re.compile(r"(?<![\w,])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\w,])")
    code_patterns = (
        ("ITEM", re.compile(r"^(?P<code>\d{6})\.\s*(?P<body>.*)$")),
        ("SUBKOMPONEN", re.compile(r"^(?P<code>\d{3}\.[A-Z0-9]{2})\s+(?P<body>.*)$", re.I)),
        ("KEGIATAN", re.compile(r"^(?P<code>[A-Z]{2,3}\.\d{4})\s*(?P<body>.*)$", re.I)),
        ("RO", re.compile(r"^(?P<code>[A-Z0-9]{3}\.[A-Z0-9]{3})\s+(?P<body>.*)$", re.I)),
        ("DETIL", re.compile(r"^(?P<code>\d{6})\s+(?P<body>.*)$")),
        ("KOMPONEN", re.compile(r"^(?P<code>\d{3})\s+(?P<body>.*)$")),
        ("PROGRAM", re.compile(r"^(?P<code>[A-Z]{2})(?P<body>Program.*)$")),
        ("KRO", re.compile(r"^(?P<code>[A-Z]{3})(?P<body>\s+.*|[A-Z].*)$")),
        ("PROGRAM", re.compile(r"^(?P<code>[A-Z]{2})\s+(?P<body>.*)$", re.I)),
    )
    parent_levels = {
        "PROGRAM": None,
        "KEGIATAN": "PROGRAM",
        "KRO": "KEGIATAN",
        "RO": "KRO",
        "KOMPONEN": "RO",
        "SUBKOMPONEN": "KOMPONEN",
        "DETIL": "SUBKOMPONEN",
        "ITEM": "DETIL",
    }
    descendants = {
        "PROGRAM": ("KEGIATAN", "KRO", "RO", "KOMPONEN", "SUBKOMPONEN", "DETIL", "ITEM"),
        "KEGIATAN": ("KRO", "RO", "KOMPONEN", "SUBKOMPONEN", "DETIL", "ITEM"),
        "KRO": ("RO", "KOMPONEN", "SUBKOMPONEN", "DETIL", "ITEM"),
        "RO": ("KOMPONEN", "SUBKOMPONEN", "DETIL", "ITEM"),
        "KOMPONEN": ("SUBKOMPONEN", "DETIL", "ITEM"),
        "SUBKOMPONEN": ("DETIL", "ITEM"),
        "DETIL": ("ITEM",),
        "ITEM": (),
    }
    current = {}
    hierarchy_rows = []
    overall = None
    pending_item = None
    wrapped_description = ""
    last_row_without_description = None

    def get_metrics(text):
        tokens = [
            match for match in number_pattern.finditer(text)
            if not re.match(r"\s*%", text[match.end():])
        ]
        if len(tokens) < 6:
            return None
        values = [to_number(match.group(0)) for match in tokens[-6:]]
        return values if all(value is not None for value in values) else None

    def make_row(kind, code, description, metrics):
        parent_level = parent_levels[kind]
        parent = current.get(parent_level) if parent_level else None
        row = [
            kind,
            code,
            description or code,
            None,
            metrics[0] if metrics else None,
            metrics[1] if metrics else None,
            metrics[4] if metrics else None,
            metrics[5] if metrics else None,
            "Laporan FA Detail",
            parent[1] if parent else None,
        ]
        hierarchy_rows.append(row)
        current[kind] = row
        for descendant in descendants[kind]:
            current.pop(descendant, None)
        return row

    for page_text in pages_text:
        for line in page_text.splitlines():
            text = line.strip()
            if not text:
                continue
            if re.search(r"\bJUMLAH SELURUHNYA\b", text, re.I):
                overall = get_metrics(text)
                continue

            matched = next(
                ((kind, match) for kind, pattern in code_patterns if (match := pattern.match(text))),
                None,
            )
            if matched:
                kind, code_match = matched
                last_row_without_description = None
                code = clean_text(code_match.group("code")).upper()
                body = code_match.group("body")
                metrics = get_metrics(body)
                number_matches = list(number_pattern.finditer(body))
                description = body[:number_matches[0].start()].strip() if number_matches else body.strip()

                if pending_item is not None:
                    pending_item = None
                if kind == "ITEM" and metrics is None:
                    parent = current.get("DETIL")
                    pending_item = {
                        "code": code,
                        "description": description,
                        "parent": parent,
                    }
                    continue
                if not metrics:
                    continue
                if kind == "ITEM" and pending_item:
                    description = pending_item["description"] or description
                if description:
                    wrapped_description = ""
                if not description and wrapped_description:
                    description = wrapped_description
                    wrapped_description = ""
                row = make_row(kind, code, description, metrics)
                if not description:
                    last_row_without_description = row
                continue

            if pending_item is not None:
                metrics = get_metrics(text)
                if metrics:
                    item = pending_item
                    parent = item["parent"]
                    row = [
                        "ITEM", item["code"], item["description"] or item["code"],
                        None, metrics[0], metrics[1], metrics[4], metrics[5],
                        "Laporan FA Detail", parent[1] if parent else None,
                    ]
                    hierarchy_rows.append(row)
                    current["ITEM"] = row
                    pending_item = None
                else:
                    pending_item["description"] = "{} {}".format(
                        pending_item["description"], text
                    ).strip()
                continue

            if (
                current
                and get_metrics(text) is None
                and not re.match(r"^\d", text)
                and not re.match(
                    r"^(?:LAPORAN REALISASI SP2D|Periode |Kementerian|Satuan Kerja|Realisasi TA|Uraian\s|Pagu Revisi|SISA\b|Hal\s)",
                    text,
                    re.I,
                )
            ):
                if last_row_without_description is not None:
                    last_row_without_description[2] = "{} {}".format(
                        "" if last_row_without_description[2] == last_row_without_description[1] else last_row_without_description[2],
                        text,
                    ).strip()
                    continue
                wrapped_description = "{} {}".format(wrapped_description, text).strip()

    rows = [BUDGET_HEADERS]
    if overall:
        rows.append([
            "TOTAL", "TOTAL", "JUMLAH SELURUHNYA", None, overall[0], None,
            overall[4], overall[5], "TOTAL", None,
        ])
    rows.extend(hierarchy_rows)
    if not overall and not hierarchy_rows:
        raise ValueError("Laporan FA terdeteksi, tetapi pagu revisi, realisasi, dan sisa tidak terbaca.")
    return rows


def normalize_ro_code(code):
    code = clean_text(code).upper()
    match = re.search(r"(?:\d{4}\.)?([A-Z0-9]{3}\.[A-Z0-9]{3})$", code)
    return match.group(1) if match else code


def budget_level_for_code(code):
    code = clean_text(code).upper()
    if re.fullmatch(r"\d{3}\.(?:[A-Z0-9]{2,3}|\d{2}\.[A-Z0-9]{2,3})", code):
        return "PROGRAM"
    if re.fullmatch(r"[A-Z]{2}", code):
        return "PROGRAM"
    if re.fullmatch(r"[A-Z]{2,3}\.\d{4}", code) or re.fullmatch(r"\d{4}", code):
        return "KEGIATAN"
    if re.fullmatch(r"\d{4}\.[A-Z0-9]{3}", code) or re.fullmatch(r"[A-Z]{3}", code):
        return "KRO"
    if re.fullmatch(r"\d{4}\.[A-Z0-9]{3}\.[A-Z0-9]{3}", code) or re.fullmatch(r"[A-Z0-9]{3}\.[A-Z0-9]{3}", code):
        return "RO"
    if re.fullmatch(r"\d{3}", code):
        return "KOMPONEN"
    if re.fullmatch(r"\d{3}\.[A-Z0-9]{2}", code) or re.fullmatch(r"[A-Z]", code):
        return "SUBKOMPONEN"
    if re.fullmatch(r"\d{6}\.", code):
        return "ITEM"
    if re.fullmatch(r"\d{6}", code):
        return "DETIL"
    return None


def normalize_hierarchy_code(level, code):
    code = clean_text(code).upper()
    if level == "PROGRAM" and "." in code:
        return code.rsplit(".", 1)[-1]
    if level == "KEGIATAN" and "." in code:
        return code.rsplit(".", 1)[-1]
    if level == "KRO" and "." in code:
        return code.rsplit(".", 1)[-1]
    if level == "RO":
        return normalize_ro_code(code)
    if level == "SUBKOMPONEN" and "." in code:
        return code.rsplit(".", 1)[-1][-1:]
    if level == "ITEM":
        return code.rstrip(".")
    return code


def summarize_revision_changes(changes):
    leaf_changes = [
        row for row in changes
        if not any(
            len(other.get("_path", ())) > len(row.get("_path", ()))
            and other.get("_path", ())[:len(row.get("_path", ()))] == row.get("_path", ())
            for other in changes
        )
    ]
    leaf_increases = [row for row in leaf_changes if row["Selisih"] > 0]
    insufficient = [row for row in leaf_increases if row["_status"] == "short"]

    def total_if_complete(key):
        values = [row.get(key) for row in leaf_changes]
        if not values or any(value is None for value in values):
            return None
        return sum(values)

    return {
        "increase_count": len(leaf_increases),
        "increase_total": sum(row["Selisih"] for row in leaf_increases),
        "decrease_total": sum(abs(row["Selisih"]) for row in leaf_changes if row["Selisih"] < 0),
        "net_change_total": sum(row["Selisih"] for row in leaf_changes),
        "pagu_semula_total": total_if_complete("Pagu semula"),
        "pagu_menjadi_total": total_if_complete("Pagu menjadi"),
        "changed_line_count": len(leaf_changes),
        "sufficient_count": sum(row["_status"] == "enough" for row in leaf_increases),
        "insufficient_count": len(insufficient),
        "unknown_count": sum(row["_status"] == "unknown" for row in leaf_increases),
        "shortfall": sum(row["Selisih"] - row["Sisa anggaran"] for row in insufficient),
    }


def normalize_budget_description(description):
    description = clean_text(description).casefold().replace("&", " dan ")
    description = re.sub(r"\[[^]]*\]|\([^)]*\)", " ", description)
    return " ".join(re.findall(r"[a-z0-9]+", description))


def budget_rows_with_paths(rows):
    paths = []
    current = {}
    for row in rows:
        kind = clean_text(row[0]).upper()
        if kind not in BUDGET_LEVEL_LABELS:
            paths.append(None)
            continue
        code = normalize_hierarchy_code(kind, row[1])
        parent_level = {
            "KEGIATAN": "PROGRAM",
            "KRO": "KEGIATAN",
            "RO": "KRO",
            "KOMPONEN": "RO",
            "SUBKOMPONEN": "KOMPONEN",
            "DETIL": "SUBKOMPONEN",
            "ITEM": "DETIL",
        }.get(kind)
        parent_path = current.get(parent_level) if parent_level else ()
        path = tuple(parent_path or ()) + ((kind, code),)
        paths.append(path)
        current[kind] = path
        level_order = list(BUDGET_LEVEL_LABELS)
        for descendant in level_order[level_order.index(kind) + 1:]:
            current.pop(descendant, None)
    return paths


def combine_budget_documents(documents):
    rkk_rows = []
    fa_rows = []
    for name, rows in documents:
        if name == "Rincian Kertas Kerja":
            rkk_rows.extend(rows[1:])
        elif name == "Laporan FA Detail":
            fa_rows.extend(rows[1:])

    rkk_total = next((row for row in rkk_rows if row[0] == "TOTAL"), None)
    fa_total = next((row for row in fa_rows if row[0] == "TOTAL"), None)
    rkk_paths = budget_rows_with_paths(rkk_rows)
    fa_paths = budget_rows_with_paths(fa_rows)
    fa_by_path = {
        path: (index, row) for index, (row, path) in enumerate(zip(fa_rows, fa_paths))
        if path is not None
    }
    fa_by_identity = {}
    for index, row in enumerate(fa_rows):
        kind = clean_text(row[0]).upper()
        if kind in BUDGET_LEVEL_LABELS:
            identity = (kind, normalize_hierarchy_code(kind, row[1]))
            fa_by_identity.setdefault(identity, []).append((index, row))
    rows = [BUDGET_HEADERS]
    rows.append([
        "TOTAL",
        "TOTAL",
        "JUMLAH SELURUHNYA",
        rkk_total[3] if rkk_total else None,
        fa_total[4] if fa_total else None,
        rkk_total[5] if rkk_total else None,
        fa_total[6] if fa_total else None,
        fa_total[7] if fa_total else None,
        "TOTAL",
        None,
    ])
    used_fa_rows = set()
    for row, path in zip(rkk_rows, rkk_paths):
        if row[0] == "TOTAL":
            continue
        row = list(row)
        kind = clean_text(row[0]).upper()
        fa_match = fa_by_path.get(path) if path is not None else None
        if fa_match and fa_match[0] in used_fa_rows:
            fa_match = None
        if fa_match is None and kind in BUDGET_LEVEL_LABELS:
            fallback = fa_by_identity.get((kind, normalize_hierarchy_code(kind, row[1])), [])
            available = [item for item in fallback if item[0] not in used_fa_rows]
            if len(available) == 1:
                fa_match = available[0]
        fa = fa_match[1] if fa_match else None
        if fa:
            row[4] = fa[4]
            row[6] = fa[6]
            row[7] = fa[7]
            row[8] = "Cocok berdasarkan kode RO" if kind == "RO" else "Cocok berdasarkan kode hierarki"
            used_fa_rows.add(fa_match[0])
        elif path is not None:
            row[8] = "Hanya di Rincian Kertas Kerja"
        rows.append(row)

    for index, (row, path) in enumerate(zip(fa_rows, fa_paths)):
        if row[0] == "TOTAL" or index in used_fa_rows:
            continue
        unmatched = list(row)
        if path is not None:
            unmatched[8] = "Hanya di Laporan FA"
        rows.append(unmatched)
    return rows


def parse_upload(name, content):
    extension = Path(name).suffix.lower()
    if len(content) > MAX_UPLOAD_SIZE:
        raise ValueError("Ukuran file melebihi batas 20 MB.")
    if extension in (".xlsx", ".xlsm", ".xls"):
        engine = "xlrd" if extension == ".xls" else "openpyxl"
        workbook = pd.ExcelFile(io.BytesIO(content), engine=engine)
        sheets = {}
        for sheet_name in workbook.sheet_names:
            frame = pd.read_excel(
                io.BytesIO(content),
                sheet_name=sheet_name,
                header=None,
                engine=engine,
                dtype=object,
            )
            sheets[sheet_name] = trim_rows(frame.where(pd.notna(frame), None).values.tolist())
        return sheets
    if extension == ".csv":
        text = None
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("CSV tidak menggunakan encoding teks yang didukung.")
        try:
            frame = pd.read_csv(io.StringIO(text), header=None, sep=None, engine="python", dtype=object)
        except pd.errors.ParserError as error:
            raise ValueError("Struktur kolom CSV tidak dapat dibaca.") from error
        return {"Data RAB": trim_rows(frame.where(pd.notna(frame), None).values.tolist())}
    if extension == ".docx":
        doc = Document(io.BytesIO(content))
        rows = []
        for table in doc.tables:
            rows.extend([[cell.text.strip() or None for cell in row.cells] for row in table.rows])
            rows.append([])
        paragraph_rows = split_text_rows("\n".join(paragraph.text for paragraph in doc.paragraphs))
        rows.extend(paragraph_rows)
        return {"Data RAB": trim_rows(rows)}
    if extension == ".rtf":
        rows = parse_rtf_rows(content)
        if not rows:
            raise ValueError("Dokumen RTF tidak memiliki teks yang dapat dibaca.")
        return {"Data RAB": rows}
    if extension == ".pdf":
        try:
            import pdfplumber
        except ImportError as error:
            raise ValueError("Dukungan PDF belum tersedia. Pasang dependensi yang tercantum di requirements.txt.") from error
        pages_text = []
        sp2d_pages_text = []
        table_rows = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables() or []:
                    table_rows.extend(table)
                pages_text.append(page.extract_text(layout=False) or "")
                sp2d_pages_text.append(page.extract_text(layout=True) or "")
        workplan = parse_rincian_kertas_text(sp2d_pages_text)
        if workplan:
            return {"Rincian Kertas Kerja": workplan}
        financial_report = parse_fa_text(sp2d_pages_text)
        if financial_report:
            return {"Laporan FA Detail": financial_report}
        sp2d = parse_sp2d_text(sp2d_pages_text)
        if sp2d:
            return {"Realisasi SP2D": sp2d}
        rows = trim_rows(table_rows) if table_rows else split_text_rows("\n".join(pages_text))
        sp2d = parse_sp2d_rows(rows)
        if sp2d:
            return {"Realisasi SP2D": sp2d}
        if not rows:
            raise ValueError("PDF tidak memiliki teks yang dapat dibaca. PDF hasil pindai perlu OCR terlebih dahulu.")
        header = find_header(rows)
        if not header:
            raise ValueError("Struktur tabel RAB pada PDF tidak dapat dipastikan. Simpan PDF dengan teks/kolom terpisah atau unggah Excel/CSV.")
        return {"Data RAB": rows}
    if extension == ".doc":
        raise ValueError("Format Word .doc lama belum didukung di Streamlit. Simpan ulang sebagai .docx lalu unggah.")
    raise ValueError("Format belum didukung. Gunakan Excel, CSV, Word .docx, atau PDF berbasis teks.")


def find_header(rows):
    best = None
    for row_index, row in enumerate(rows[:40]):
        labels = [clean_text(value) for value in row]
        normalized = [label.lower() for label in labels]
        desc_index = next(
            (
                index for index, label in enumerate(normalized)
                if re.fullmatch(
                    r"uraian(?: kegiatan)?|deskripsi(?: kegiatan)?|nama kegiatan|rincian(?: kegiatan)?|kegiatan|belanja",
                    label,
                )
            ),
            -1,
        )
        if desc_index < 0:
            continue
        amount_matches = sum(bool(re.search(r"jumlah|pagu|anggaran|nilai|sisa", label)) for label in normalized)
        score = amount_matches * 2 + sum(bool(label) for label in normalized)
        if best is None or score >= best["score"]:
            best = {"row_index": row_index, "labels": labels, "normalized": normalized, "description": desc_index, "score": score}
    return best


def find_revision(rows):
    layout = None
    for marker_index, row in enumerate(rows[:40]):
        markers = [clean_text(value).lower() for value in row]
        original_start = markers.index("semula") if "semula" in markers else -1
        revised_start = markers.index("menjadi") if "menjadi" in markers else -1
        if original_start < 0 or revised_start <= original_start:
            continue
        for header_index in range(marker_index + 1, min(len(rows), marker_index + 5)):
            labels = [clean_text(value).lower() for value in rows[header_index]]
            original_desc = next((i for i, label in enumerate(labels) if original_start <= i < revised_start and label == "uraian"), -1)
            revised_desc = next((i for i, label in enumerate(labels) if i >= revised_start and label == "uraian"), -1)
            delta_col = next((i for i, label in enumerate(labels) if i >= revised_start and label == "selisih"), -1)
            remaining_col = next((i for i, label in enumerate(labels) if label == "sisa anggaran"), -1)
            orig_amounts = [
                i for i, label in enumerate(labels)
                if original_start <= i < revised_start and re.sub(r"\s+", " ", label).strip() == "jumlah"
            ]
            original_code = next(
                (i for i, label in enumerate(labels) if original_start <= i < revised_start and label == "kode anggaran"),
                -1,
            )
            revised_code = next(
                (i for i, label in enumerate(labels) if i >= revised_start and label == "kode anggaran"),
                -1,
            )
            revised_amount = next(
                (i for i, label in enumerate(labels) if revised_start <= i < delta_col and label == "jumlah"),
                -1,
            ) if delta_col >= 0 else -1
            if original_desc >= 0 and revised_desc >= 0 and delta_col >= 0 and orig_amounts and revised_amount >= 0:
                original_amount = max(
                    orig_amounts,
                    key=lambda column: sum(to_number(row[column] if column < len(row) else None) is not None for row in rows[header_index + 1:]),
                )
                layout = {
                    "first": header_index + 1,
                    "original_code": original_code,
                    "original_desc": original_desc,
                    "original_amount": original_amount,
                    "revised_code": revised_code,
                    "revised_desc": revised_desc,
                    "revised_amount": revised_amount,
                    "delta": delta_col,
                    "remaining": remaining_col,
                }
                break
        if layout:
            break
    if not layout:
        return None

    changes = []
    current_hierarchy = {}
    level_order = list(BUDGET_LEVEL_LABELS)
    for row_index in range(layout["first"], len(rows)):
        row = rows[row_index]

        def get(column):
            return row[column] if column >= 0 and column < len(row) else None

        original_desc = clean_text(get(layout["original_desc"]))
        revised_desc = clean_text(get(layout["revised_desc"]))
        original_code = clean_text(get(layout["original_code"]))
        revised_code = clean_text(get(layout["revised_code"]))
        code = revised_code or original_code
        level = budget_level_for_code(code) if code else None
        original_amount = to_number(get(layout["original_amount"]))
        revised_amount = to_number(get(layout["revised_amount"]))
        if original_amount is None and not original_code and not original_desc and revised_desc:
            original_amount = 0
        if revised_amount is None and not revised_code and not revised_desc and original_desc:
            revised_amount = 0
        if code and level is None:
            original_code = ""
            revised_code = ""
            code = ""
        if level:
            current_hierarchy[level] = normalize_hierarchy_code(level, code)
            for descendant in level_order[level_order.index(level) + 1:]:
                current_hierarchy.pop(descendant, None)
        if not code and (not original_desc and not revised_desc):
            continue
        delta = to_number(get(layout["delta"]))
        calculated_delta = (
            revised_amount - original_amount
            if original_amount is not None and revised_amount is not None
            else None
        )
        if delta is None or (delta == 0 and calculated_delta not in (None, 0)):
            delta = calculated_delta
        if delta is None or delta == 0:
            continue
        remaining = to_number(get(layout["remaining"]))
        status = "decrease" if delta < 0 else "unknown" if remaining is None else "enough" if remaining > delta else "short"
        change_level = BUDGET_LEVEL_LABELS.get(level, "Rincian item" if current_hierarchy else "Rincian")
        hierarchy_path = tuple((item, current_hierarchy[item]) for item in level_order if item in current_hierarchy)
        if not code and revised_desc:
            hierarchy_path += (("ITEM", re.sub(r"\s+", " ", revised_desc).casefold()),)
        changes.append({
            "Uraian semula": original_desc or "—",
            "Kode semula": original_code or "—",
            "Pagu semula": original_amount,
            "Uraian menjadi": revised_desc or "—",
            "Kode menjadi": revised_code or "—",
            "Pagu menjadi": revised_amount,
            "Selisih": delta,
            "Sisa anggaran": remaining,
            "Sisa setelah kenaikan": remaining - delta if delta > 0 and remaining is not None else None,
            "Tingkat": change_level,
            "_path": hierarchy_path,
            "Kecukupan": {
                "decrease": "Pengurangan",
                "unknown": "Cek sisa",
                "enough": "Sisa cukup",
                "short": "Sisa tidak lebih besar",
            }[status],
            "_status": status,
        })
    increases = [row for row in changes if row["Selisih"] > 0]
    decreases = [row for row in changes if row["Selisih"] < 0]
    return {
        "changes": changes,
        "increases": increases,
        "decreases": decreases,
        **summarize_revision_changes(changes),
    }


def analyze_budget_rows(rows, sheet_name):
    header = rows[0]
    positions = {name: header.index(name) for name in BUDGET_HEADERS}
    total = next((row for row in rows[1:] if row[positions["_Jenis"]] == "TOTAL"), [])
    prepared = []
    for index, raw in enumerate(rows[1:], start=2):
        kind = clean_text(raw[positions["_Jenis"]])
        if kind == "TOTAL":
            continue
        code = clean_text(raw[positions["Kode Anggaran"]])
        description = clean_text(raw[positions["Uraian"]])
        if not code and not description:
            continue
        full = to_number(raw[positions["Pagu Penuh"]])
        effective = to_number(raw[positions["Pagu Efektif"]])
        prepared.append({
            "Tingkat": BUDGET_LEVEL_LABELS.get(kind, kind.title()),
            "Kode": code,
            "Induk": clean_text(raw[positions["Induk"]]),
            "Uraian": description,
            "Pagu": effective if effective is not None else full,
            "Pagu Penuh": full,
            "Pagu Efektif": effective,
            "Pagu Terblokir": to_number(raw[positions["Pagu Terblokir"]]),
            "Realisasi": to_number(raw[positions["Realisasi"]]),
            "Sisa anggaran": to_number(raw[positions["Sisa Anggaran"]]),
            "Status Pasangan": clean_text(raw[positions["Status Pasangan"]]),
            "Baris": index,
            "_raw": list(raw),
        })

    def total_value(name):
        return to_number(total[positions[name]]) if total else None

    comparison = {
        "pagu_penuh": total_value("Pagu Penuh"),
        "pagu_efektif": total_value("Pagu Efektif"),
        "pagu_terblokir": total_value("Pagu Terblokir"),
        "realisasi": total_value("Realisasi"),
        "sisa": total_value("Sisa Anggaran"),
        "matched_count": sum(
            row["Tingkat"] == "RO" and row["Status Pasangan"] == "Cocok berdasarkan kode RO"
            for row in prepared
        ),
        "unmatched_count": sum(
            row["Tingkat"] == "RO" and row["Status Pasangan"] != "Cocok berdasarkan kode RO"
            for row in prepared
        ),
        "hierarchy_matched_count": sum(
            row["Status Pasangan"] == "Cocok berdasarkan kode hierarki"
            or row["Status Pasangan"] == "Cocok berdasarkan kode RO"
            for row in prepared
        ),
        "hierarchy_unmatched_count": sum(
            row["Status Pasangan"].startswith("Hanya di")
            for row in prepared
        ),
    }
    hierarchy_counts = {
        kind: sum(row["Tingkat"] == label for row in prepared)
        for kind, label in BUDGET_LEVEL_LABELS.items()
    }
    category_levels = {}
    for kind, label in BUDGET_LEVEL_LABELS.items():
        categories = [
            {
                "Tingkat": label,
                "Kode": row["Kode"],
                "Kategori": "{}{}{}".format(
                    "{} · ".format(row["Kode"]) if row["Kode"] else "",
                    row["Uraian"],
                    " (induk {})".format(row["Induk"]) if row["Induk"] else "",
                ),
                "Pagu": row["Pagu"],
                "Pagu Penuh": row["Pagu Penuh"],
                "Pagu Efektif": row["Pagu Efektif"],
                "Pagu Terblokir": row["Pagu Terblokir"],
                "Realisasi": row["Realisasi"],
                "Sisa": row["Sisa anggaran"],
            }
            for row in prepared
            if row["Tingkat"] == label and row["Pagu"] is not None and row["Pagu"] > 0
        ]
        if categories:
            categories.sort(key=lambda row: row["Pagu"], reverse=True)
            category_levels[kind] = categories
    category_levels["ALL"] = sorted(
        [
            dict(category, Kategori="{} · {}".format(category["Tingkat"], category["Kategori"]))
            for categories in category_levels.values()
            for category in categories
        ],
        key=lambda row: row["Pagu"],
        reverse=True,
    )
    default_level = next(
        (kind for kind in ("AKUN", "DETIL", "RO", "KEGIATAN", "PROGRAM") if kind in category_levels),
        next(iter(category_levels), None),
    )
    categories = category_levels.get(default_level, [])
    hierarchy_summary = ", ".join(
        "{} {}".format(count, BUDGET_LEVEL_LABELS[kind])
        for kind, count in hierarchy_counts.items() if count
    )
    insights = []
    if hierarchy_summary:
        insights.append("Hierarki yang dibaca: {}.".format(hierarchy_summary))
    if comparison["pagu_terblokir"] is not None:
        insights.append("Pagu terblokir yang ditandai * pada Rincian Kertas Kerja berjumlah {}.".format(
            format_rupiah(comparison["pagu_terblokir"])
        ))
    if comparison["pagu_penuh"] is not None and comparison["pagu_efektif"] is not None:
        difference = comparison["pagu_penuh"] - comparison["pagu_efektif"]
        insights.append("Selisih pagu penuh dan pagu efektif: {}.".format(format_rupiah(difference)))
    if comparison["pagu_efektif"] and comparison["realisasi"] is not None:
        insights.append("Realisasi mencapai {:.1f}% dari pagu efektif.".format(
            comparison["realisasi"] / comparison["pagu_efektif"] * 100
        ))
    if comparison["unmatched_count"]:
        insights.append("{} kode RO belum memiliki pasangan pada kedua file; periksa status pasangan di tabel.".format(
            comparison["unmatched_count"]
        ))
    if categories:
        top = categories[0]
        insights.append("{} memiliki pagu terbesar pada tingkat {}.".format(
            top["Kategori"], BUDGET_LEVEL_LABELS[default_level].lower()
        ))

    return {
        "sheet": sheet_name,
        "header_found": True,
        "revision": None,
        "rows": prepared,
        "ranked": sorted(prepared, key=lambda row: row["Pagu"] or 0, reverse=True)[:5],
        "categories": categories,
        "category_levels": category_levels,
        "default_category_level": default_level,
        "hierarchy_counts": hierarchy_counts,
        "columns": [(index, name) for index, name in enumerate(header) if name and not name.startswith("_")],
        "money_col": positions["Pagu Efektif"],
        "remaining_col": positions["Sisa Anggaran"],
        "pagu_total": comparison["pagu_efektif"] or comparison["pagu_penuh"],
        "remaining_total": comparison["sisa"],
        "category_total": sum(row["Pagu"] for row in categories),
        "category_remaining_total": sum(row["Sisa"] for row in categories if row["Sisa"] is not None),
        "grouping": BUDGET_LEVEL_LABELS.get(default_level, "rincian anggaran"),
        "source_count": len(prepared),
        "excluded_count": 0,
        "insights": insights[:5],
        "budget_comparison": comparison,
    }


def analyze_rows(rows, sheet_name):
    if rows and rows[0] and rows[0][0] == "_Jenis":
        return analyze_budget_rows(rows, sheet_name)
    header = find_header(rows)
    revision = find_revision(rows)
    if not header:
        return {
            "sheet": sheet_name,
            "header_found": False,
            "revision": revision,
            "rows": [],
            "categories": [],
            "columns": [],
            "pagu_total": None,
            "remaining_total": None,
            "source_count": len(rows),
            "excluded_count": len(rows),
            "grouping": "",
            "insights": ["Struktur tabel belum dikenali. Pastikan ada kolom uraian/kegiatan dan angka anggaran."],
        }

    header_row = header["row_index"]
    labels = header["labels"]
    normalized = header["normalized"]
    start = 0
    for previous_row in rows[max(0, header_row - 3):header_row]:
        marker = next((i for i, value in enumerate(previous_row) if clean_text(value).lower() == "menjadi"), -1)
        if marker >= 0:
            start = marker

    columns = [(i, label) for i, label in enumerate(labels) if i >= start and label]
    description_col = next((i for i, label in columns if re.search(r"uraian|deskripsi|kegiatan|rincian|belanja", label.lower())), header["description"])
    code_col = next((i for i, label in columns if re.search(r"kode anggaran|^kode$|^akun$", label.lower())), -1)
    amount_candidates = [(i, label) for i, label in columns if re.search(r"jumlah|pagu|total anggaran|nilai anggaran", label.lower())]
    money_col = amount_candidates[-1][0] if amount_candidates else -1
    remaining_candidates = [(i, label) for i, label in columns if re.search(r"sisa anggaran|^sisa$", label.lower())]
    remaining_col = remaining_candidates[-1][0] if remaining_candidates else -1
    prepared = []
    source_count = max(0, len(rows) - header_row - 1)
    for row_index, raw in enumerate(rows[header_row + 1:], header_row + 2):
        row = list(raw)
        description = clean_text(row[description_col] if description_col < len(row) else None)
        code = clean_text(row[code_col] if code_col >= 0 and code_col < len(row) else None)
        amount = to_number(row[money_col] if money_col >= 0 and money_col < len(row) else None)
        remaining = to_number(row[remaining_col] if remaining_col >= 0 and remaining_col < len(row) else None)
        if not description or re.fullmatch(r"\d+(?:[.,]\d+)?", description):
            continue
        if description.lower() in ("uraian", "kode anggaran", "vol.", "satuan", "harga satuan", "jumlah", "pagu", "sisa anggaran"):
            continue
        if amount is None and remaining is None:
            continue
        prepared.append({"Kode": code, "Uraian": description, "Pagu": amount, "Sisa anggaran": remaining, "Baris": row_index, "_raw": row})

    hierarchy_current = {}
    for row in prepared:
        kind = budget_level_for_code(row["Kode"])
        if kind == "SUBKOMPONEN" and "KOMPONEN" not in hierarchy_current:
            kind = None
        if kind:
            parent_level = {
                "KEGIATAN": "PROGRAM",
                "KRO": "KEGIATAN",
                "RO": "KRO",
                "KOMPONEN": "RO",
                "SUBKOMPONEN": "KOMPONEN",
                "DETIL": "SUBKOMPONEN",
                "ITEM": "DETIL",
            }.get(kind)
            parent_code = hierarchy_current.get(parent_level) if parent_level else None
            row.update({
                "Tingkat": BUDGET_LEVEL_LABELS[kind],
                "Induk": parent_code or "",
                "Pagu Penuh": row["Pagu"],
                "Pagu Efektif": None,
                "Pagu Terblokir": None,
                "Realisasi": None,
            })
            hierarchy_current[kind] = normalize_hierarchy_code(kind, row["Kode"])
            for descendant in list(BUDGET_LEVEL_LABELS)[list(BUDGET_LEVEL_LABELS).index(kind) + 1:]:
                hierarchy_current.pop(descendant, None)
        elif (
            hierarchy_current.get("DETIL")
            and row["Uraian"].startswith(("-", "•"))
        ):
            row.update({
                "Tingkat": BUDGET_LEVEL_LABELS["ITEM"],
                "Induk": hierarchy_current["DETIL"],
                "Pagu Penuh": row["Pagu"],
                "Pagu Efektif": None,
                "Pagu Terblokir": None,
                "Realisasi": None,
            })

    positive = [row["Pagu"] for row in prepared if row["Pagu"] is not None and row["Pagu"] > 0]
    pagu_total = max(positive) if positive else None
    detailed = [
        row for row in prepared
        if row["Pagu"] is not None and row["Pagu"] > 0
        and (pagu_total is None or row["Pagu"] < pagu_total * 0.9)
        and (row["Uraian"].startswith(("-", "•")) or row["Kode"])
    ]
    ranked = sorted(detailed or [row for row in prepared if row["Pagu"] and row["Pagu"] > 0], key=lambda row: row["Pagu"], reverse=True)[:5]
    activity = [row for row in prepared if re.fullmatch(r"\d{3}", row["Kode"]) and row["Pagu"] and row["Pagu"] > 0]
    account = [row for row in prepared if re.fullmatch(r"\d{6}", row["Kode"]) and row["Pagu"] and row["Pagu"] > 0]
    categories = []
    grouping = "kelompok kegiatan"
    if len(activity) >= 2:
        categories = [{"Kode": row["Kode"], "Kategori": row["Kode"] + " · " + row["Uraian"], "Pagu": row["Pagu"], "Sisa": row["Sisa anggaran"]} for row in activity]
    elif account:
        grouped = {}
        for row in account:
            category = grouped.setdefault(row["Kode"], {"Kode": row["Kode"], "Kategori": row["Kode"] + " · " + row["Uraian"], "Pagu": 0, "Sisa": 0, "_sisa_count": 0, "_count": 0})
            category["Pagu"] += row["Pagu"]
            category["_count"] += 1
            if row["Sisa anggaran"] is not None:
                category["Sisa"] += row["Sisa anggaran"]
                category["_sisa_count"] += 1
        categories = [
            {key: value for key, value in category.items() if not key.startswith("_")}
            for category in grouped.values()
        ]
        for category, source in zip(categories, grouped.values()):
            if source["_sisa_count"] != source["_count"]:
                category["Sisa"] = None
        grouping = "akun belanja"
    else:
        categories = [{"Kode": row["Kode"], "Kategori": row["Uraian"], "Pagu": row["Pagu"], "Sisa": row["Sisa anggaran"]} for row in detailed if row["Pagu"] and row["Pagu"] > 0]
        grouping = "rincian belanja"

    category_levels = {}
    for kind, label in BUDGET_LEVEL_LABELS.items():
        level_categories = [
            {
                "Tingkat": label,
                "Kode": row["Kode"],
                "Kategori": "{}{}{}".format(
                    "{} · ".format(row["Kode"]) if row["Kode"] else "",
                    row["Uraian"],
                    " (induk {})".format(row["Induk"]) if row.get("Induk") else "",
                ),
                "Pagu": row["Pagu"],
                "Pagu Penuh": row.get("Pagu Penuh"),
                "Pagu Efektif": row.get("Pagu Efektif"),
                "Pagu Terblokir": row.get("Pagu Terblokir"),
                "Realisasi": row.get("Realisasi"),
                "Sisa": row["Sisa anggaran"],
            }
            for row in prepared
            if row.get("Tingkat") == label and row["Pagu"] is not None and row["Pagu"] > 0
        ]
        if level_categories:
            category_levels[kind] = level_categories
    default_level = next(
        (
            kind for kind in (
                "DETIL", "ITEM", "RO", "KEGIATAN", "KRO",
                "SUBKOMPONEN", "KOMPONEN", "PROGRAM",
            ) if kind in category_levels
        ),
        None,
    )
    if category_levels:
        category_levels["ALL"] = sorted(
            [
                dict(category, Kategori="{} · {}".format(category["Tingkat"], category["Kategori"]))
                for level, items in category_levels.items()
                for category in items
            ],
            key=lambda row: row["Pagu"],
            reverse=True,
        )
        categories = category_levels[default_level]
        grouping = BUDGET_LEVEL_LABELS[default_level]
    categories.sort(key=lambda row: row["Pagu"], reverse=True)
    category_total = sum(row["Pagu"] for row in categories)
    remaining_categories = [row for row in categories if row["Sisa"] is not None]
    category_remaining_total = sum(row["Sisa"] for row in remaining_categories) if remaining_categories else None

    insights = []
    if revision and revision["increases"]:
        if revision["insufficient_count"]:
            insights.append("{} baris kenaikan tingkat terendah memiliki sisa lebih kecil dari selisih; kekurangan terhitung {}.".format(
                revision["insufficient_count"], format_rupiah(revision["shortfall"])
            ))
        elif revision["unknown_count"]:
            insights.append("{} baris kenaikan tingkat terendah, {} belum memiliki nilai sisa yang bisa dibandingkan.".format(
                revision.get("increase_count", len(revision["increases"])), revision["unknown_count"]
            ))
        else:
            insights.append("Sisa pada semua {} baris kenaikan tingkat terendah yang terbaca mencukupi.".format(
                revision.get("increase_count", len(revision["increases"]))
            ))
    if categories and category_total:
        top = categories[0]
        insights.append("{} merupakan porsi pagu terbesar ({:.0f}%).".format(top["Kategori"], top["Pagu"] / category_total * 100))
        if category_remaining_total is not None:
            insights.append("Sisa yang terbaca pada kategori terpilih berjumlah {}.".format(format_rupiah(category_remaining_total)))
        if any(row["Sisa"] is not None and (row["Sisa"] < 0 or row["Sisa"] > row["Pagu"]) for row in categories):
            insights.append("Ada nilai sisa di luar rentang pagu; periksa perubahan anggaran atau formula sumber.")
    elif not categories:
        insights.append("Belum ada kategori bernilai yang bisa dipetakan untuk grafik.")

    return {
        "sheet": sheet_name,
        "header_found": True,
        "revision": revision,
        "rows": prepared,
        "ranked": ranked,
        "categories": categories,
        "category_levels": category_levels,
        "default_category_level": default_level,
        "hierarchy_counts": {
            kind: sum(row.get("Tingkat") == label for row in prepared)
            for kind, label in BUDGET_LEVEL_LABELS.items()
        },
        "columns": columns,
        "money_col": money_col,
        "remaining_col": remaining_col,
        "pagu_total": pagu_total,
        "remaining_total": max([row["Sisa anggaran"] for row in prepared if row["Sisa anggaran"] is not None], default=None),
        "category_total": category_total,
        "category_remaining_total": category_remaining_total,
        "grouping": grouping,
        "source_count": source_count,
        "excluded_count": source_count - len(prepared),
        "insights": insights[:4],
    }


def analyze_combined_tables(tables):
    analyses = [(name, analyze_rows(rows, name)) for name, rows in tables]
    valid = [(name, analysis) for name, analysis in analyses if analysis.get("header_found")]
    combined_rows = []
    row_indexes = {}
    duplicate_sources = {}
    conflicts = {}

    for source, analysis in valid:
        for source_row in analysis["rows"]:
            row = dict(source_row)
            row.pop("_raw", None)
            row["Sumber"] = source
            amount_values = tuple(
                row.get(key)
                for key in ("Pagu", "Pagu Penuh", "Pagu Efektif", "Pagu Terblokir", "Realisasi", "Sisa anggaran")
            )
            code = clean_text(row.get("Kode")).upper()
            description = clean_text(row.get("Uraian")).casefold()
            identity = (
                code, description, amount_values
            ) if code else (source, description, amount_values)
            group_identity = (code, description) if code else None

            if group_identity:
                conflicts.setdefault(group_identity, set()).add(row.get("Pagu"))
            if identity in row_indexes:
                existing = combined_rows[row_indexes[identity]]
                duplicate_sources.setdefault(identity, set()).add(existing["Sumber"])
                duplicate_sources[identity].add(source)
                continue
            row_indexes[identity] = len(combined_rows)
            combined_rows.append(row)

    for identity, sources in duplicate_sources.items():
        combined_rows[row_indexes[identity]]["Sumber"] = ", ".join(sorted(sources))

    conflict_groups = [
        identity for identity, values in conflicts.items()
        if len([value for value in values if value is not None]) > 1
    ]

    financial_sources = {
        source for source, analysis in valid if analysis.get("budget_comparison")
    }
    if financial_sources:
        parent_levels = {
            "KEGIATAN": "PROGRAM",
            "KRO": "KEGIATAN",
            "RO": "KRO",
            "KOMPONEN": "RO",
            "SUBKOMPONEN": "KOMPONEN",
            "DETIL": "SUBKOMPONEN",
            "ITEM": "DETIL",
        }
        grouped_rows = {}
        non_hierarchy_rows = []
        for row in combined_rows:
            kind = next(
                (level for level, label in BUDGET_LEVEL_LABELS.items() if label == row.get("Tingkat")),
                None,
            )
            if not kind or not row.get("Kode"):
                non_hierarchy_rows.append(row)
                continue
            parent_level = parent_levels.get(kind)
            parent_code = (
                normalize_hierarchy_code(parent_level, row.get("Induk", ""))
                if parent_level and row.get("Induk") else ""
            )
            identity = (
                kind,
                normalize_hierarchy_code(kind, row["Kode"]),
                parent_code,
            )
            grouped_rows.setdefault(identity, []).append(row)

        deduplicated_rows = list(non_hierarchy_rows)
        for group in grouped_rows.values():
            financial_rows = [
                row for row in group
                if financial_sources.intersection(row.get("Sumber", "").split(", "))
            ]
            if not financial_rows:
                deduplicated_rows.extend(group)
                continue
            primary = dict(financial_rows[0])
            for row in group:
                if row is primary or row in financial_rows:
                    continue
                for key in (
                    "Pagu Penuh", "Pagu Terblokir", "Realisasi", "Sisa anggaran",
                ):
                    if primary.get(key) is None and row.get(key) is not None:
                        primary[key] = row[key]
                if (
                    primary.get("Uraian") in ("", primary.get("Kode"))
                    and row.get("Uraian")
                ):
                    primary["Uraian"] = row["Uraian"]
                sources = set(primary.get("Sumber", "").split(", "))
                sources.update(row.get("Sumber", "").split(", "))
                primary["Sumber"] = ", ".join(sorted(source for source in sources if source))
            deduplicated_rows.append(primary)
        combined_rows = deduplicated_rows

    category_indexes = {}
    categories = []
    for source, analysis in valid:
        for source_category in analysis.get("categories", []):
            category = dict(source_category)
            category["Sumber"] = source
            identity = (
                clean_text(category.get("Kode")).upper(),
                clean_text(category.get("Kategori")).casefold(),
                category.get("Pagu"),
                category.get("Sisa"),
            )
            if identity in category_indexes:
                existing = categories[category_indexes[identity]]
                if source not in existing["Sumber"].split(", "):
                    existing["Sumber"] += ", " + source
                continue
            category_indexes[identity] = len(categories)
            categories.append(category)
    categories.sort(key=lambda row: row["Pagu"] or 0, reverse=True)

    category_levels = {}
    for kind, label in BUDGET_LEVEL_LABELS.items():
        level_categories = [
            {
                "Tingkat": label,
                "Kode": row["Kode"],
                "Kategori": "{}{}{}".format(
                    "{} · ".format(row["Kode"]) if row["Kode"] else "",
                    row["Uraian"],
                    " (induk {})".format(row["Induk"]) if row.get("Induk") else "",
                ),
                "Pagu": row["Pagu"],
                "Pagu Penuh": row.get("Pagu Penuh"),
                "Pagu Efektif": row.get("Pagu Efektif"),
                "Pagu Terblokir": row.get("Pagu Terblokir"),
                "Realisasi": row.get("Realisasi"),
                "Sisa": row.get("Sisa anggaran"),
            }
            for row in combined_rows
            if row.get("Tingkat") == label and row.get("Pagu") is not None and row["Pagu"] > 0
        ]
        if level_categories:
            category_levels[kind] = level_categories
    default_level = next(
        (
            kind for kind in (
                "DETIL", "ITEM", "RO", "KEGIATAN", "KRO",
                "SUBKOMPONEN", "KOMPONEN", "PROGRAM",
            ) if kind in category_levels
        ),
        None,
    )
    if category_levels:
        category_levels["ALL"] = sorted(
            [
                dict(category, Kategori="{} · {}".format(category["Tingkat"], category["Kategori"]))
                for items in category_levels.values()
                for category in items
            ],
            key=lambda row: row["Pagu"],
            reverse=True,
        )
        categories = category_levels[default_level]

    revisions = [
        (source, analysis["revision"])
        for source, analysis in valid if analysis.get("revision")
    ]
    revision = None
    if revisions:
        changes = []
        for source, source_revision in revisions:
            for change in source_revision["changes"]:
                changes.append(dict(change, Sumber=source))
        increases = [row for row in changes if row["Selisih"] > 0]
        revision = {
            "changes": changes,
            "increases": increases,
            "decreases": [row for row in changes if row["Selisih"] < 0],
        }
        fa_analyses = [
            (source, analysis) for source, analysis in valid
            if analysis.get("budget_comparison")
        ]
        apply_fa_remaining_to_revision(revision, fa_analyses)

    insights = []
    for source, analysis in valid:
        insights.extend("{}: {}".format(source, insight) for insight in analysis.get("insights", []))
    if conflict_groups:
        insights.append(
            "{} kode dan uraian memiliki nilai pagu berbeda antar sumber; tinjau baris bertanda sumber sebelum menjumlahkannya.".format(
                len(conflict_groups)
            )
        )
    if duplicate_sources:
        insights.append(
            "{} baris yang identik dari beberapa file disatukan agar tidak dihitung berulang.".format(
                len(duplicate_sources)
            )
        )
    if len(valid) > 1 and combined_rows:
        insights.append(
            "Analisis gabungan memuat {} baris bernilai dari {} tabel/file; kolom Sumber menunjukkan asal datanya.".format(
                len(combined_rows), len(valid)
            )
        )

    comparisons = [
        analysis.get("budget_comparison")
        for _, analysis in valid if analysis.get("budget_comparison")
    ]
    comparison = comparisons[0] if len(comparisons) == 1 else None
    total_pagu = sum(row["Pagu"] for row in categories if row.get("Pagu") is not None)
    remaining_values = [row["Sisa"] for row in categories if row.get("Sisa") is not None]
    sources = list(dict.fromkeys(name for name, _ in valid))
    return {
        "sheet": "Gabungan semua file",
        "header_found": bool(valid),
        "revision": revision,
        "rows": combined_rows,
        "ranked": sorted(
            combined_rows,
            key=lambda row: row.get("Pagu") or 0,
            reverse=True,
        )[:5],
        "categories": categories,
        "category_levels": category_levels,
        "default_category_level": default_level,
        "hierarchy_counts": {
            kind: sum(row.get("Tingkat") == label for row in combined_rows)
            for kind, label in BUDGET_LEVEL_LABELS.items()
        },
        "columns": [],
        "money_col": -1,
        "remaining_col": -1,
        "pagu_total": total_pagu or None,
        "remaining_total": sum(remaining_values) if remaining_values else None,
        "category_total": total_pagu,
        "category_remaining_total": sum(remaining_values) if remaining_values else None,
        "grouping": "gabungan lintas sumber",
        "source_count": sum(analysis.get("source_count", 0) for _, analysis in valid),
        "excluded_count": sum(analysis.get("excluded_count", 0) for _, analysis in analyses),
        "sources": sources,
        "insights": insights[:12],
        "budget_comparison": comparison,
    }


def apply_fa_remaining_to_revision(revision, fa_analyses):
    fa_rows = []
    fa_paths = []
    for source, analysis in fa_analyses:
        source_rows = analysis.get("rows", [])
        raw_rows = [row.get("_raw") for row in source_rows if row.get("_raw")]
        paths = budget_rows_with_paths(raw_rows)
        for row, path in zip(source_rows, paths):
            if row.get("Tingkat") in BUDGET_LEVEL_LABELS.values():
                fa_rows.append(dict(row, Sumber=source))
                fa_paths.append(path)

    for change in revision["changes"]:
        kind = next(
            (level for level, label in BUDGET_LEVEL_LABELS.items() if label == change["Tingkat"]),
            None,
        )
        code = change.get("Kode menjadi", "—")
        if code == "—":
            code = change.get("Kode semula", "—")
        code_key = normalize_hierarchy_code(kind, code) if kind and code != "—" else None
        candidates = []
        for row, path in zip(fa_rows, fa_paths):
            row_kind = next(
                (level for level, label in BUDGET_LEVEL_LABELS.items() if label == row["Tingkat"]),
                None,
            )
            if kind and row_kind == kind and code_key is not None:
                if normalize_hierarchy_code(kind, row["Kode"]) == code_key:
                    candidates.append((row, path))
            elif not code_key and row_kind == "ITEM":
                candidates.append((row, path))

        change_path = change.get("_path", ())
        matching_basis = "Kode dan hierarki"
        if kind == "ITEM":
            parent_path = change_path[:-1]
            item_candidates = [
                (row, path)
                for row, path in zip(fa_rows, fa_paths)
                if row.get("Tingkat") == BUDGET_LEVEL_LABELS["ITEM"]
                and path
                and path[:-1] == parent_path
            ]
            code_matches = [
                item for item in candidates
                if code_key
                and normalize_hierarchy_code("ITEM", item[0].get("Kode", "")) == code_key
                and item[1]
                and item[1][:-1] == parent_path
            ]
            candidates = code_matches or item_candidates
            if code_matches:
                matching_basis = "Kode dan hierarki"
            normalized_description = normalize_budget_description(change["Uraian menjadi"])
            exact_description = [
                item for item in candidates
                if normalize_budget_description(item[0]["Uraian"]) == normalized_description
            ]
            if exact_description and not code_matches:
                candidates = exact_description
                matching_basis = "Uraian rincian cocok"
            elif candidates and not code_matches:
                scored = sorted(
                    [
                        (
                            SequenceMatcher(
                                None,
                                normalized_description,
                                normalize_budget_description(item[0]["Uraian"]),
                            ).ratio(),
                            item,
                        )
                        for item in candidates
                    ],
                    key=lambda match: match[0],
                    reverse=True,
                )
                if (
                    scored[0][0] >= 0.82
                    and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.08)
                ):
                    candidates = [scored[0][1]]
                    matching_basis = "Uraian rincian paling mirip"
                else:
                    candidates = []
        else:
            exact_path = [item for item in candidates if item[1] == change_path]
            if exact_path:
                candidates = exact_path
            if len(candidates) > 1:
                description = normalize_budget_description(change["Uraian menjadi"])
                matching_description = [
                    item for item in candidates
                    if normalize_budget_description(item[0]["Uraian"]) == description
                ]
                if matching_description:
                    candidates = matching_description

        remaining_values = [item[0].get("Sisa anggaran") for item in candidates]
        remaining = (
            sum(remaining_values)
            if remaining_values and all(value is not None for value in remaining_values)
            else None
        )
        if kind == "ITEM" and remaining is None:
            parent_path = change_path[:-1]
            for prefix_length in range(len(parent_path), 0, -1):
                ancestor_path = parent_path[:prefix_length]
                ancestor_candidates = [
                    (row, path)
                    for row, path in zip(fa_rows, fa_paths)
                    if path == ancestor_path and row.get("Sisa anggaran") is not None
                ]
                if len(ancestor_candidates) == 1:
                    candidates = ancestor_candidates
                    remaining = ancestor_candidates[0][0]["Sisa anggaran"]
                    ancestor_kind = ancestor_path[-1][0]
                    matching_basis = "Sisa FA pada {}".format(
                        BUDGET_LEVEL_LABELS[ancestor_kind].lower()
                    )
                    break
        if len(candidates) > 1 and matching_basis == "Uraian rincian cocok":
            matching_basis = "Gabungan rincian beruraian sama"
        change["Sisa anggaran"] = remaining
        change["Basis pasangan FA"] = matching_basis if remaining is not None else "Sisa FA tidak ditemukan"
        change["Sumber Sisa FA"] = ", ".join(sorted({
            row.get("Sumber", "")
            for row, _ in candidates
            if row.get("Sumber")
        }))
        change["Sisa setelah kenaikan"] = (
            remaining - change["Selisih"]
            if change["Selisih"] > 0 and remaining is not None
            else None
        )
        status = (
            "decrease" if change["Selisih"] < 0
            else "unknown" if remaining is None
            else "enough" if remaining > change["Selisih"]
            else "short"
        )
        change["_status"] = status
        change["Kecukupan"] = {
            "decrease": "Pengurangan",
            "enough": "Sisa cukup",
            "short": "Sisa tidak lebih besar",
            "unknown": "Sisa FA tidak ditemukan",
        }[status]

    changes = revision["changes"]
    revision.update({
        "increases": [row for row in changes if row["Selisih"] > 0],
        "decreases": [row for row in changes if row["Selisih"] < 0],
        **summarize_revision_changes(changes),
    })


def format_rupiah(value):
    if value is None:
        return "—"
    return "Rp {:,.0f}".format(value).replace(",", ".")


def format_compact_rupiah(value):
    if value is None:
        return "—"
    amount = abs(value)
    if amount >= 1_000_000_000:
        scaled, suffix = value / 1_000_000_000, " M"
    elif amount >= 1_000_000:
        scaled, suffix = value / 1_000_000, " jt"
    elif amount >= 1_000:
        scaled, suffix = value / 1_000, " rb"
    else:
        return "Rp {:,.0f}".format(value).replace(",", ".")
    number = "{:,.1f}".format(scaled).replace(",", "X").replace(".", ",").replace("X", ".")
    if number.endswith(",0"):
        number = number[:-2]
    return number + suffix


def load_analysis():
    uploads = st.file_uploader(
        "Unggah file anggaran",
        type=["xlsx", "xls", "xlsm", "csv", "docx", "rtf", "pdf"],
        help="Pilih satu atau beberapa file Excel, CSV, Word .docx/.rtf, atau PDF berbasis teks. Beberapa file akan dianalisis bersama dan tetap tersedia per sumber. Maksimum 20 MB per file; PDF hasil pindai perlu OCR.",
        accept_multiple_files=True,
    )
    if uploads:
        try:
            file_contents = [(uploaded.name, uploaded.getvalue()) for uploaded in uploads]
            key_digest = hashlib.sha256()
            for filename, content in file_contents:
                key_digest.update(filename.encode("utf-8"))
                key_digest.update(content)
            upload_key = key_digest.hexdigest()
            if upload_key != st.session_state.get("upload_key"):
                parsed_files = [
                    (filename, parse_upload(filename, content))
                    for filename, content in file_contents
                ]
                special_documents = []
                source_tables = []
                for filename, sheets in parsed_files:
                    special_sheets = [
                        (sheet_name, rows)
                        for sheet_name, rows in sheets.items()
                        if sheet_name in ("Rincian Kertas Kerja", "Laporan FA Detail")
                    ]
                    if special_sheets:
                        special_documents.extend(special_sheets)
                    else:
                        for sheet_name, rows in sheets.items():
                            label = "{} — {}".format(Path(filename).stem, sheet_name)
                            source_tables.append((label, rows))
                if special_documents:
                    if len(special_documents) > 2:
                        raise ValueError("Unggah maksimal satu Rincian Kertas Kerja dan satu Laporan FA Detail.")
                    document_types = [name for name, _ in special_documents]
                    if len(set(document_types)) != len(document_types):
                        raise ValueError("Pilih satu Rincian Kertas Kerja dan satu Laporan FA Detail, bukan dua file dengan jenis yang sama.")
                    source_tables.insert(0, (
                        "Perbandingan Anggaran",
                        combine_budget_documents(special_documents),
                    ))
                if not source_tables or not any(
                    any(clean_text(value) for row in rows for value in row)
                    for _, rows in source_tables
                ):
                    raise ValueError("Tidak ada isi tabel yang dapat dibaca dari file.")
                sheets = {}
                analyses = {}
                for label, rows in source_tables:
                    unique_label = label
                    suffix = 2
                    while unique_label in sheets:
                        unique_label = "{} ({})".format(label, suffix)
                        suffix += 1
                    sheets[unique_label] = rows
                    analyses[unique_label] = analyze_rows(rows, unique_label)
                if len(source_tables) > 1:
                    combined_label = "Gabungan semua file"
                    analyses[combined_label] = analyze_combined_tables(source_tables)
                    default_selection = combined_label
                else:
                    default_selection = next(iter(sheets))
                st.session_state.upload_name = ", ".join(filename for filename, _ in file_contents)
                st.session_state.upload_key = upload_key
                st.session_state.sheets = sheets
                st.session_state.analyses = analyses
                st.session_state.selected_sheet = default_selection
                st.session_state.pop("tor_result", None)
                for key in list(st.session_state):
                    if key.startswith((
                        "tor_beneficiaries_",
                        "tor_strategy_",
                        "tor_schedule_",
                        "tor_kpi_output_",
                        "tor_kpi_outcome_",
                        "tor_cost_benefit_",
                        "tor_risks_",
                    )):
                        st.session_state.pop(key, None)
                for key in ("tor_activity", "tor_location", "tor_year", "tor_notes", "tor_ro_code", "_tor_activity_ro"):
                    st.session_state.pop(key, None)
        except Exception as error:
            st.error("File tidak dapat diproses: {}".format(error))
            return None
    sheets = st.session_state.get("sheets")
    if not sheets:
        st.info("Unggah RAB untuk melihat ringkasan, grafik, dan membuat draft TOR.")
        return None
    names = list(sheets)
    chosen = st.selectbox("Pilih sheet / tabel", names, index=names.index(st.session_state.get("selected_sheet", names[0])) if st.session_state.get("selected_sheet") in names else 0)
    st.session_state.selected_sheet = chosen
    analysis = st.session_state.get("analyses", {}).get(chosen)
    if analysis is None:
        analysis = analyze_rows(sheets[chosen], chosen)
    return chosen, analysis


def hierarchy_budget_totals(analysis, selected_level):
    comparison = analysis.get("budget_comparison")
    if not comparison:
        return None
    if selected_level == "ALL":
        return comparison

    categories = analysis.get("category_levels", {}).get(selected_level, [])
    selected_label = BUDGET_LEVEL_LABELS.get(selected_level)
    selected_codes = {category["Kode"] for category in categories}
    rows = analysis.get("rows", [])
    selected_rows = [
        row for row in rows
        if row.get("Tingkat") == selected_label and row.get("Kode") in selected_codes
    ]

    child_level = {
        "Program": "Kegiatan",
        "Kegiatan": "KRO",
        "KRO": "RO",
        "RO": "Komponen",
        "Komponen": "Subkomponen",
        "Subkomponen": "Detil belanja",
        "Detil belanja": "Rincian item",
    }
    rows_by_level = {}
    for row in rows:
        rows_by_level.setdefault(row.get("Tingkat"), []).append(row)

    def descendant_values(parent, key):
        child_label = child_level.get(parent["Tingkat"])
        children = [
            row for row in rows_by_level.get(child_label, [])
            if row.get("Induk") == parent["Kode"]
        ] if child_label else []
        if not children:
            return []
        if key == "Pagu Penuh":
            values = []
            for child in children:
                value = child.get(key)
                if value is None:
                    continue
                if value != 0:
                    values.append(value)
                    continue
                nested_values = descendant_values(child, key)
                if nested_values:
                    values.extend(nested_values)
                elif value is not None:
                    values.append(value)
            return values
        values = []
        for child in children:
            if child.get(key) is not None:
                values.append(child[key])
            else:
                nested_values = descendant_values(child, key)
                values.extend(nested_values)
        return values

    def aggregate_level(key):
        if not selected_rows:
            return None
        values = []
        for row in selected_rows:
            if key == "Pagu Penuh" and row.get(key) == 0:
                descendants = descendant_values(row, key)
                if descendants:
                    values.extend(descendants)
                else:
                    values.append(0)
                continue
            if key == "Pagu Penuh" and row.get(key) is None:
                continue
            if row.get(key) is not None:
                values.append(row[key])
            else:
                descendants = descendant_values(row, key)
                values.extend(descendants)
        return sum(values) if values else None

    blocked_values = [row.get("Pagu Terblokir") for row in selected_rows]
    totals = {
        "pagu_penuh": aggregate_level("Pagu Penuh"),
        "pagu_terblokir": (
            sum(value for value in blocked_values if value is not None)
            if any(value is not None for value in blocked_values) else None
        ),
        "pagu_efektif": aggregate_level("Pagu Efektif"),
        "realisasi": aggregate_level("Realisasi"),
        "sisa": aggregate_level("Sisa anggaran"),
    }
    return totals


def render_dashboard(data):
    if not data:
        return
    sheet, analysis = data
    st.subheader("Ringkasan anggaran")
    st.caption("{} · {} baris anggaran disiapkan; {} baris disaring.".format(
        analysis["sheet"], len(analysis["rows"]), analysis["excluded_count"]
    ))
    if analysis.get("sources"):
        st.caption("Sumber gabungan: {}".format(" · ".join(analysis["sources"])))
    comparison = analysis.get("budget_comparison")
    category_levels = analysis.get("category_levels", {})
    if category_levels:
        level_options = ["ALL"] + [level for level in category_levels if level != "ALL"]
        default_level = analysis.get("default_category_level")
        level_index = level_options.index(default_level) if default_level in level_options else 0
        selected_level = st.selectbox(
            "Tingkat hierarki untuk grafik dan kartu anggaran",
            level_options,
            index=level_index,
            format_func=lambda level: "Semua level" if level == "ALL" else BUDGET_LEVEL_LABELS.get(level, level.title()),
            key="budget_hierarchy_level",
        )
        categories = category_levels[selected_level]
    else:
        selected_level = None
        categories = analysis.get("categories", [])
    has_fa_comparison = bool(comparison and comparison.get("pagu_efektif") is not None)
    if comparison:
        selected_totals = hierarchy_budget_totals(analysis, selected_level)
        level_label = (
            "semua level" if selected_level == "ALL"
            else BUDGET_LEVEL_LABELS.get(selected_level, "tingkat terpilih")
        )
        metric_columns = st.columns(5)
        metric_columns[0].metric("Pagu penuh · {}".format(level_label), format_compact_rupiah(selected_totals["pagu_penuh"]))
        metric_columns[1].metric("Pagu efektif · {}".format(level_label), format_compact_rupiah(selected_totals["pagu_efektif"]))
        metric_columns[2].metric("Pagu terblokir (*) · {}".format(level_label), format_compact_rupiah(selected_totals["pagu_terblokir"]))
        metric_columns[3].metric("Realisasi · {}".format(level_label), format_compact_rupiah(selected_totals["realisasi"]))
        metric_columns[4].metric("Sisa anggaran · {}".format(level_label), format_compact_rupiah(selected_totals["sisa"]))
        if selected_level == "ALL":
            st.caption("Semua level dipilih. Kartu menampilkan total satker, bukan penjumlahan baris lintas hierarki.")
        else:
            st.caption(
                "Kartu menjumlahkan nilai yang tersedia pada tingkat {}; baris tanpa data/pasangan tidak masuk hitungan. "
                "Pagu penuh pada induk yang bernilai 0 dihitung dari turunan terdekat.".format(level_label.lower())
            )
        st.caption(
            "{} tingkat hierarki cocok; {} tingkat belum memiliki pasangan. "
            "Dari RO: {} cocok, {} belum cocok.".format(
                comparison.get("hierarchy_matched_count", comparison["matched_count"]),
                comparison.get("hierarchy_unmatched_count", comparison["unmatched_count"]),
                comparison["matched_count"],
                comparison["unmatched_count"],
            )
        )
    else:
        if selected_level == "ALL":
            pagu_for_card = analysis.get("pagu_total")
            remaining_for_card = analysis.get("remaining_total")
        else:
            pagu_for_card = sum(row["Pagu"] for row in categories) if categories else analysis.get("pagu_total")
            remaining_values = [row["Sisa"] for row in categories if row["Sisa"] is not None]
            remaining_for_card = sum(remaining_values) if categories else analysis.get("remaining_total")
        col1, col2, col3, col4 = st.columns(4)
        category_label = (
            "semua level" if selected_level == "ALL"
            else BUDGET_LEVEL_LABELS.get(selected_level, "kategori")
        )
        col1.metric("Total pagu · {}".format(category_label), format_compact_rupiah(pagu_for_card))
        col2.metric("Sisa anggaran · {}".format(category_label), format_compact_rupiah(remaining_for_card))
        col3.metric("Baris pada tingkat terpilih", len(categories))
        col4.metric("Rincian bernilai", len(analysis["rows"]))

    has_fa_source = any(
        "fa detail" in source.casefold() or "perbandingan anggaran" in source.casefold()
        for source in analysis.get("sources", [])
    )
    revision = analysis.get("revision")
    if revision:
        st.markdown("### Perubahan RAB & kecukupan dana")
        if has_fa_source:
            st.caption(
                "Pagu semula dan menjadi diambil dari Excel revisi; Sisa FA dipasangkan dari laporan FA Detail "
                "berdasarkan kode/jalur hierarki dan uraian rincian."
            )
        if revision["insufficient_count"]:
            st.warning("{} kenaikan pada tingkat terendah perlu ditinjau: sisa FA tidak lebih besar dari selisih.".format(revision["insufficient_count"]))
        elif revision["unknown_count"]:
            st.info("{} kenaikan pada tingkat terendah belum dapat dibandingkan karena nilai sisa tidak ditemukan.".format(revision["unknown_count"]))
        elif revision["increases"]:
            st.success("Sisa anggaran pada semua kenaikan tingkat terendah yang terbaca mencukupi.")
        else:
            st.info("Tidak ada kenaikan pagu yang terdeteksi.")
        if revision.get("pagu_semula_total") is None or revision.get("pagu_menjadi_total") is None:
            st.info("Total pagu semula/menjadi pada baris perubahan belum lengkap; periksa nilai pada tabel.")
        metric_columns = st.columns(5)
        metric_columns[0].metric(
            "Pagu semula · baris berubah",
            format_compact_rupiah(revision.get("pagu_semula_total")),
        )
        metric_columns[1].metric(
            "Pagu menjadi · baris berubah",
            format_compact_rupiah(revision.get("pagu_menjadi_total")),
        )
        metric_columns[2].metric(
            "Selisih bersih",
            format_compact_rupiah(revision.get("net_change_total")),
        )
        metric_columns[3].metric(
            "Total kenaikan",
            format_compact_rupiah(revision["increase_total"]),
        )
        metric_columns[4].metric(
            "Total penurunan",
            format_compact_rupiah(revision["decrease_total"]),
        )
        rev1, rev2, rev3, rev4 = st.columns(4)
        rev1.metric("Sisa mencukupi", revision["sufficient_count"])
        rev2.metric("Sisa tidak lebih besar", revision["insufficient_count"])
        rev3.metric("Sisa FA tidak ditemukan", revision["unknown_count"])
        rev4.metric("Kekurangan terhitung", format_compact_rupiah(revision["shortfall"]))
        st.caption(
            "Ringkasan pagu hanya memakai baris perubahan pada tingkat terendah agar perubahan induk-anak tidak dihitung ganda. "
            "Kenaikan dinilai cukup hanya jika sisa FA lebih besar dari selisih (sisa setelah kenaikan > Rp0)."
        )
        if revision["changes"]:
            revision_frame = pd.DataFrame(revision["changes"]).drop(columns=["_status", "_path"], errors="ignore")
            st.dataframe(revision_frame, use_container_width=True, hide_index=True)

    if categories:
        chart_frame = pd.DataFrame(categories)
        chart_frame["Pagu"] = pd.to_numeric(chart_frame["Pagu"], errors="coerce")
        chart_frame["Sisa"] = pd.to_numeric(chart_frame["Sisa"], errors="coerce")
        if selected_level == "ALL":
            st.info("Grafik menampilkan baris dari semua tingkat; nilai induk dan anak tidak dijumlahkan menjadi total.")
        st.markdown("### {} terbesar per {}".format(
            "Pagu efektif" if has_fa_comparison and selected_level == "RO" else "Pagu",
            "semua level" if selected_level == "ALL" else BUDGET_LEVEL_LABELS.get(selected_level, "kategori").lower() if selected_level else "kategori",
        ))
        pagu_data = chart_frame.nlargest(10, "Pagu").sort_values("Pagu")
        axis_divisor = 1_000_000_000 if pagu_data["Pagu"].max() >= 1_000_000_000 else 1_000_000 if pagu_data["Pagu"].max() >= 1_000_000 else 1
        axis_unit = "miliar" if axis_divisor == 1_000_000_000 else "juta" if axis_divisor == 1_000_000 else "rupiah"
        axis_suffix = " M" if axis_divisor == 1_000_000_000 else " jt" if axis_divisor == 1_000_000 else ""
        pagu_data = pagu_data.assign(
            **{
                "Pagu (skala)": pagu_data["Pagu"] / axis_divisor,
                "Label ringkas": pagu_data["Pagu"].map(format_compact_rupiah),
            }
        )
        fig = px.bar(
            pagu_data,
            x="Pagu (skala)",
            y="Kategori",
            orientation="h",
            text="Label ringkas",
            color_discrete_sequence=[COLORS[0]],
            custom_data=["Pagu"],
        )
        fig.update_traces(
            textposition="outside",
            cliponaxis=False,
            hovertemplate="%{y}<br>Pagu: Rp %{customdata[0]:,.0f}<extra></extra>",
        )
        fig.update_layout(
            height=max(330, 48 * len(pagu_data)),
            xaxis_title="Pagu (Rp {})".format(axis_unit),
            separators=",.",
            xaxis=dict(ticksuffix=axis_suffix),
            yaxis_title=None,
            showlegend=False,
            margin=dict(l=10, r=55, t=10, b=40),
        )
        st.plotly_chart(fig, use_container_width=True)

        if has_fa_comparison and selected_level == "RO":
            compare_columns = [
                column for column in ("Pagu Penuh", "Pagu Efektif", "Realisasi")
                if chart_frame[column].notna().any()
            ]
            if len(compare_columns) > 1:
                comparison_data = chart_frame.nlargest(8, "Pagu").copy()
                comparison_long = comparison_data.melt(
                    id_vars=["Kategori"],
                    value_vars=compare_columns,
                    var_name="Jenis anggaran",
                    value_name="Nilai",
                ).dropna(subset=["Nilai"])
                chart_divisor = 1_000_000_000 if comparison_long["Nilai"].max() >= 1_000_000_000 else 1_000_000
                chart_unit = "miliar" if chart_divisor == 1_000_000_000 else "juta"
                comparison_long["Nilai (skala)"] = comparison_long["Nilai"] / chart_divisor
                comparison_fig = px.bar(
                    comparison_long,
                    x="Nilai (skala)",
                    y="Kategori",
                    color="Jenis anggaran",
                    orientation="h",
                    barmode="group",
                    color_discrete_sequence=COLORS[:len(compare_columns)],
                    custom_data=["Nilai"],
                )
                comparison_fig.update_traces(
                    hovertemplate="%{y}<br>%{fullData.name}: Rp %{customdata[0]:,.0f}<extra></extra>"
                )
                comparison_fig.update_layout(
                    height=max(360, 54 * min(len(comparison_data), 8)),
                    xaxis_title="Nilai (Rp {})".format(chart_unit),
                    separators=",.",
                    yaxis_title=None,
                    margin=dict(l=10, r=25, t=10, b=40),
                )
                st.markdown("### Perbandingan pagu dan realisasi per RO")
                st.plotly_chart(comparison_fig, use_container_width=True)

        if selected_level != "ALL":
            st.markdown("### Porsi {} ".format("pagu efektif" if has_fa_comparison and selected_level == "RO" else "pagu"))
            pie_data = chart_frame.nlargest(5, "Pagu").copy()
            pie_data["Label"] = pie_data.apply(
                lambda row: row["Kode"] or row["Kategori"][:24],
                axis=1,
            )
            if len(chart_frame) > 5:
                other = chart_frame.nsmallest(len(chart_frame) - 5, "Pagu")["Pagu"].sum()
                pie_data = pd.concat(
                    [
                        pie_data,
                        pd.DataFrame([{"Kategori": "Kategori lainnya", "Label": "Lainnya", "Pagu": other}]),
                    ],
                    ignore_index=True,
                )
            pie = px.pie(
                pie_data,
                values="Pagu",
                names="Label",
                hole=0.58,
                color_discrete_sequence=COLORS,
                custom_data=["Kategori"],
            )
            pie.update_traces(
                textposition="inside",
                textinfo="percent",
                hovertemplate="%{customdata[0]}<br>Rp %{value:,.0f} (%{percent})<extra></extra>",
                sort=True,
            )
            pie.update_layout(
                height=410,
                margin=dict(l=10, r=240, t=10, b=10),
                legend_title=None,
                legend=dict(orientation="v", yanchor="middle", y=0.5, xanchor="left", x=0.68, font=dict(size=12)),
            )
            st.plotly_chart(pie, use_container_width=True)
            remaining = chart_frame.dropna(subset=["Sisa"]).copy()
            if not remaining.empty:
                remaining = remaining.nlargest(8, "Sisa")
                st.markdown("### Sisa anggaran terbesar")
                remaining = remaining.sort_values("Sisa")
                remaining_divisor = 1_000_000_000 if remaining["Sisa"].max() >= 1_000_000_000 else 1_000_000 if remaining["Sisa"].max() >= 1_000_000 else 1
                remaining_unit = "miliar" if remaining_divisor == 1_000_000_000 else "juta" if remaining_divisor == 1_000_000 else "rupiah"
                remaining_suffix = " M" if remaining_divisor == 1_000_000_000 else " jt" if remaining_divisor == 1_000_000 else ""
                remaining = remaining.assign(
                    **{
                        "Sisa (skala)": remaining["Sisa"] / remaining_divisor,
                        "Label ringkas": remaining["Sisa"].map(format_compact_rupiah),
                    }
                )
                bar = px.bar(
                    remaining,
                    x="Sisa (skala)",
                    y="Kategori",
                    orientation="h",
                    text="Label ringkas",
                    color_discrete_sequence=[COLORS[1]],
                    custom_data=["Sisa"],
                )
                bar.update_traces(
                    textposition="outside",
                    cliponaxis=False,
                    hovertemplate="%{y}<br>Sisa: Rp %{customdata[0]:,.0f}<extra></extra>",
                )
                bar.update_layout(
                    height=max(280, min(440, 48 * len(remaining))),
                    xaxis_title="Sisa (Rp {})".format(remaining_unit),
                    separators=",.",
                    xaxis=dict(ticksuffix=remaining_suffix),
                    yaxis_title=None,
                    showlegend=False,
                    margin=dict(l=10, r=45, t=10, b=40),
                )
                st.plotly_chart(bar, use_container_width=True)
    st.markdown("### Temuan dari data")
    for insight in analysis.get("insights", []):
        st.markdown("- " + insight)

    st.markdown("### Data anggaran yang sudah diproses")
    search = st.text_input("Cari uraian atau kode")
    prepared = pd.DataFrame(analysis["rows"])
    if not prepared.empty:
        if search:
            mask = prepared["Uraian"].str.contains(search, case=False, na=False) | prepared["Kode"].str.contains(search, case=False, na=False)
            prepared = prepared.loc[mask]
        st.dataframe(prepared.drop(columns=["_raw"], errors="ignore"), use_container_width=True, hide_index=True)
    else:
        st.warning("Struktur tabel belum dikenali. Pastikan terdapat header uraian dan pagu/jumlah.")


def infer_activity(analysis):
    coded = next((
        row for row in analysis.get("rows", [])
        if row.get("Tingkat") == "RO" and row["Pagu"] is not None
    ), None)
    if coded:
        return coded["Uraian"]
    if analysis.get("ranked"):
        return re.sub(r"^[\-•]\s*", "", analysis["ranked"][0]["Uraian"])
    return analysis.get("rows", [{}])[0].get("Uraian", "")


def build_tor_sections(activity, location, year, notes, analysis, selected_row=None):
    rows = analysis.get("rows", [])
    source = selected_row or next(
        (row for row in rows if row["Uraian"].lower() == activity.lower() and row["Pagu"] is not None),
        None,
    )
    if source:
        if analysis.get("budget_comparison"):
            amount_text = "Kode RO: {}. Pagu penuh: {}. Pagu efektif: {}. Pagu terblokir: {}. Realisasi: {}. Sisa anggaran: {}.".format(
                source["Kode"] or "—",
                format_rupiah(source.get("Pagu Penuh")),
                format_rupiah(source.get("Pagu Efektif")),
                format_rupiah(source.get("Pagu Terblokir")),
                format_rupiah(source.get("Realisasi")),
                format_rupiah(source.get("Sisa anggaran")),
            )
        else:
            amount_text = "Pagu: {}{}".format(
                format_rupiah(source["Pagu"]),
                " · kode {}".format(source["Kode"]) if source["Kode"] else "",
            )
            amount_text += " · Sisa anggaran: {}".format(format_rupiah(source["Sisa anggaran"]))
    else:
        amount_text = "Baris kegiatan belum ditemukan. Cocokkan nama kegiatan dengan RAB sebelum menetapkan pagu."
    purpose = notes or "Lengkapi tujuan khusus dan indikator keberhasilan kegiatan."
    return [
        ("I. Latar Belakang", "TOR ini menjadi acuan pelaksanaan {} pada tahun {}, berdasarkan data pada sheet {}.".format(activity, year, analysis.get("sheet", "RAB"))),
        ("II. Maksud dan Tujuan", purpose),
        ("III. Sasaran", "Sasaran dan penerima manfaat: lengkapi sesuai target kegiatan."),
        ("IV. Ruang Lingkup", "Persiapan, koordinasi, pelaksanaan, pemantauan, dan pelaporan; sesuaikan dengan kebutuhan kegiatan."),
        ("V. Keluaran", "Target keluaran terukur: lengkapi sebelum TOR disahkan."),
        ("VI. Waktu dan Tempat", "Tahun: {} · Lokasi: {}. Lengkapi jadwal pelaksanaan secara rinci.".format(year, location)),
        ("VII. Pembiayaan", amount_text),
        ("VIII. Penutup", "Lengkapi bagian sasaran, keluaran, jadwal, dan persetujuan. Verifikasi pembiayaan dengan RAB yang disahkan."),
    ]


def find_local_tor_template():
    templates = sorted(ROOT.glob("[[]Template[]]*.docx"))
    return templates[0] if templates else None


def set_docx_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def docx_table_frame(table, first_data_row=1, column_names=None):
    if column_names is None:
        column_names = [clean_text(cell.text) for cell in table.rows[first_data_row - 1].cells]
    rows = [
        [cell.text for cell in row.cells]
        for row in table.rows[first_data_row:]
    ]
    return pd.DataFrame(rows, columns=column_names)


def docx_frame_rows(frame):
    return [
        ["" if pd.isna(value) else str(value) for value in row]
        for row in frame.astype(object).values.tolist()
    ]


def tor_budget_amount(budget_data):
    amount = budget_data.get("Pagu Efektif")
    return amount if amount is not None else budget_data.get("Pagu Penuh")


def write_docx_table_rows(table, rows, first_data_row=1, remove_extra=False):
    available = len(table.rows) - first_data_row
    if len(rows) > available:
        raise ValueError("Jumlah baris penyesuaian melebihi struktur tabel pada template TOR.")
    for row_index, values in enumerate(rows, start=first_data_row):
        table_row = table.rows[row_index]
        if len(values) != len(table_row.cells):
            raise ValueError("Jumlah kolom penyesuaian tidak sesuai dengan tabel template TOR.")
        for cell, value in zip(table_row.cells, values):
            set_docx_paragraph_text(cell.paragraphs[0], str(value))
            for paragraph in cell.paragraphs[1:]:
                set_docx_paragraph_text(paragraph, "")
    if remove_extra:
        for row in list(table.rows[first_data_row + len(rows):]):
            table._tbl.remove(row._tr)


def get_tor_hierarchy(analysis, selected_row):
    rows = analysis.get("rows", [])
    selected_index = next(
        (index for index, row in enumerate(rows) if row is selected_row),
        -1,
    )
    if selected_index < 0:
        return {}, []

    hierarchy = {"RO": selected_row}
    child = selected_row
    child_index = selected_index
    parent_levels = {"RO": "KRO", "KRO": "KEGIATAN", "KEGIATAN": "PROGRAM"}
    for parent_level in parent_levels.values():
        parent_code = child.get("Induk", "")
        if not parent_code:
            break
        parent = next(
            (
                (index, row) for index, row in reversed(list(enumerate(rows[:child_index])))
                if row.get("Kode") == parent_code and row.get("Tingkat") == BUDGET_LEVEL_LABELS[parent_level]
            ),
            None,
        )
        if parent is None:
            break
        child_index, child = parent
        hierarchy[parent_level] = child

    descendants = []
    for row in rows[selected_index + 1:]:
        if row.get("Tingkat") in ("Program", "Kegiatan", "KRO", "RO"):
            break
        if row.get("Tingkat") in ("Komponen", "Subkomponen", "Detil belanja"):
            descendants.append(row)
    return hierarchy, descendants


def tor_strategy_default(selected_row, descendants):
    lines = ["Keluaran yang dituju: {}.".format(selected_row["Uraian"])]
    if descendants:
        lines.append("Rincian kegiatan berdasarkan Rincian Kertas Kerja:")
        lines.extend(
            "- {} {} — {}".format(row["Tingkat"], row["Kode"], row["Uraian"]).strip()
            for row in descendants
        )
    else:
        lines.append("Rincian pelaksanaan disesuaikan dengan komponen kegiatan pada Rincian Kertas Kerja.")
    return "\n".join(lines)


def replace_tor_strategy(document, strategy_text):
    paragraphs = document.paragraphs
    method_index = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "Metode Pelaksanaan"),
        None,
    )
    end_index = next(
        (
            index for index, paragraph in enumerate(paragraphs)
            if method_index is not None and index > method_index
            and paragraph.text.strip() == "Tahapan dan Waktu Pelaksanaan"
        ),
        None,
    )
    if method_index is None or end_index is None:
        raise ValueError("Bagian strategi pelaksanaan tidak ditemukan pada template TOR.")

    target_paragraphs = paragraphs[method_index + 1:end_index]
    lines = [line.strip() for line in strategy_text.splitlines() if line.strip()]
    if not target_paragraphs:
        raise ValueError("Template TOR tidak menyediakan paragraf untuk strategi pelaksanaan.")
    if not lines:
        lines = [""]
    if len(lines) > len(target_paragraphs):
        lines = lines[:len(target_paragraphs) - 1] + ["\n".join(lines[len(target_paragraphs) - 1:])]

    for paragraph, line in zip(target_paragraphs, lines):
        set_docx_paragraph_text(paragraph, line)
    for paragraph in target_paragraphs[len(lines):]:
        paragraph._element.getparent().remove(paragraph._element)


def set_tor_budget_cost(document, budget_data):
    amount = tor_budget_amount(budget_data)
    if amount is None:
        replacement = "[Pagu belum tersedia pada dashboard]"
    else:
        replacement = format_rupiah(amount)

    paragraph = next(
        (
            paragraph for paragraph in document.paragraphs
            if paragraph.text.startswith("Biaya yang diperlukan untuk kegiatan ini sebesar")
        ),
        None,
    )
    if paragraph is None:
        raise ValueError("Bagian Biaya yang Diperlukan tidak ditemukan pada template TOR.")
    runs = paragraph.runs
    for run_index, run in enumerate(runs):
        match = re.match(r"(\s*)([\d.,]+)(.*)", run.text)
        if match and any(
            previous.text.rstrip().endswith("Rp")
            for previous in runs[:run_index]
        ):
            run.text = match.group(1) + replacement.replace("Rp ", "") + match.group(3)
            return
    text = paragraph.text
    updated, count = re.subn(r"(?<=Rp )[\d.,]+", replacement.replace("Rp ", ""), text, count=1)
    if not count:
        raise ValueError("Nilai biaya pada template TOR tidak dapat diperbarui.")
    set_docx_paragraph_text(paragraph, updated)


def make_tor_from_local_template(ro_code, budget_data, hierarchy, adjustments):
    template_path = find_local_tor_template()
    if template_path is None:
        raise FileNotFoundError("Template TOR Word tidak ditemukan di folder aplikasi.")

    document = Document(str(template_path))
    if len(document.tables) < 10:
        raise ValueError("Struktur tabel template TOR tidak sesuai dengan template yang didukung.")

    metadata = document.tables[0]
    for row_index, level in ((0, "PROGRAM"), (1, "KEGIATAN"), (2, "KRO"), (3, "RO")):
        item = hierarchy.get(level)
        code = item["Kode"] if item else "[Tidak tersedia di dashboard]"
        description = item["Uraian"] if item else "[Tidak tersedia di dashboard]"
        if level == "RO":
            code = ro_code or code
            description = budget_data.get("Uraian", description)
        set_docx_paragraph_text(metadata.cell(row_index, 1).paragraphs[0], code)
        set_docx_paragraph_text(metadata.cell(row_index, 2).paragraphs[0], description)

    for key, table_index, first_data_row, remove_extra in (
        ("beneficiaries", 4, 1, False),
        ("schedule", 5, 2, False),
        ("kpi_output", 6, 1, False),
        ("kpi_outcome", 7, 1, False),
        ("cost_benefit", 8, 1, False),
        ("risks", 9, 1, False),
    ):
        if key in adjustments:
            write_docx_table_rows(
                document.tables[table_index],
                adjustments[key],
                first_data_row=first_data_row,
                remove_extra=remove_extra,
            )
    if "strategy" in adjustments:
        replace_tor_strategy(document, adjustments["strategy"])
    set_tor_budget_cost(document, budget_data)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_tor_docx(ro_code, budget_data, hierarchy, adjustments):
    if not find_local_tor_template():
        raise FileNotFoundError("Template TOR Word tidak ditemukan di folder aplikasi.")
    if not budget_data:
        raise ValueError("Pilih satu baris RO dari dashboard sebelum membuat TOR.")
    return make_tor_from_local_template(
        ro_code,
        budget_data,
        hierarchy,
        adjustments,
    )


def render_tor(data):
    st.subheader("Generator TOR")
    if not data:
        st.info("Unggah RAB pada tab Dashboard untuk menghubungkan data dan membuat draft TOR.")
        return
    _, analysis = data
    template_path = find_local_tor_template()
    if template_path is None:
        st.error("Template TOR Word tidak ditemukan. Pastikan file [Template] TOR tersedia di folder aplikasi.")
        return

    ro_rows = [
        row for row in analysis["rows"]
        if row.get("Tingkat") == "RO" and row["Kode"]
    ]
    if not ro_rows:
        st.warning("Belum ada baris Rincian Output (RO) yang dapat digunakan untuk membuat TOR.")
        return

    row_by_code = {row["Kode"]: row for row in ro_rows}
    selected_code = st.selectbox(
        "Pilih kode Rincian Output (RO)",
        list(row_by_code),
        key="tor_ro_code",
        format_func=lambda code: "{} — {}".format(code, row_by_code[code]["Uraian"]),
    )
    selected_row = row_by_code[selected_code]
    hierarchy, descendants = get_tor_hierarchy(analysis, selected_row)
    template = Document(str(template_path))
    if len(template.tables) < 10:
        st.error("Struktur tabel template TOR tidak sesuai dengan template yang didukung.")
        return

    identity = hashlib.sha1(selected_code.encode("utf-8")).hexdigest()[:10]
    st.caption(
        "TOR akan menggunakan format dan isi dokumen template. Program, Kegiatan, KRO, dan RO "
        "diisi dari hierarki dashboard."
    )
    st.markdown("#### Identitas kegiatan dari dashboard")
    hierarchy_columns = st.columns(4)
    for column, level in zip(hierarchy_columns, ("PROGRAM", "KEGIATAN", "KRO", "RO")):
        item = hierarchy.get(level)
        column.markdown("**{}**".format(BUDGET_LEVEL_LABELS[level]))
        column.write(
            "{} — {}".format(item["Kode"], item["Uraian"])
            if item else "Tidak tersedia pada data anggaran"
        )

    amount = tor_budget_amount(selected_row)
    st.metric("Biaya yang diperlukan dari dashboard", format_rupiah(amount))
    st.caption("Sumber biaya: pagu efektif; jika tidak tersedia, memakai pagu penuh.")

    beneficiaries = docx_table_frame(template.tables[4])
    schedule_columns = ["No", "Rincian Aktivitas"] + [
        "Bulan {}".format(month) for month in range(1, 13)
    ]
    schedule = docx_table_frame(template.tables[5], first_data_row=2, column_names=schedule_columns)
    kpi_output = docx_table_frame(template.tables[6])
    kpi_outcome = docx_table_frame(template.tables[7])
    cost_benefit = docx_table_frame(template.tables[8])
    risks = docx_table_frame(template.tables[9])

    with st.expander("Sesuaikan bagian TOR yang dapat diubah", expanded=True):
        st.markdown("**Penerima Manfaat**")
        beneficiaries = st.data_editor(
            beneficiaries,
            use_container_width=True,
            hide_index=True,
            key="tor_beneficiaries_{}".format(identity),
        )
        st.markdown("**Strategi Pencapaian Keluaran**")
        strategy = st.text_area(
            "Uraian strategi dan rincian sesuai RO",
            value=tor_strategy_default(selected_row, descendants),
            height=240,
            key="tor_strategy_{}".format(identity),
        )
        st.markdown("**Kurun Waktu Pencapaian Keluaran**")
        schedule = st.data_editor(
            schedule,
            use_container_width=True,
            hide_index=True,
            height=420,
            key="tor_schedule_{}".format(identity),
        )
        st.markdown("**KPI Output**")
        kpi_output = st.data_editor(
            kpi_output,
            use_container_width=True,
            hide_index=True,
            key="tor_kpi_output_{}".format(identity),
        )
        st.markdown("**KPI Outcome**")
        kpi_outcome = st.data_editor(
            kpi_outcome,
            use_container_width=True,
            hide_index=True,
            key="tor_kpi_outcome_{}".format(identity),
        )
        st.markdown("**Analisis Biaya dan Manfaat**")
        cost_benefit = st.data_editor(
            cost_benefit,
            use_container_width=True,
            hide_index=True,
            key="tor_cost_benefit_{}".format(identity),
        )
        st.markdown("**Manajemen Risiko**")
        risks = st.data_editor(
            risks,
            use_container_width=True,
            hide_index=True,
            key="tor_risks_{}".format(identity),
        )

    adjustments = {
        "beneficiaries": docx_frame_rows(beneficiaries),
        "strategy": strategy,
        "schedule": docx_frame_rows(schedule),
        "kpi_output": docx_frame_rows(kpi_output),
        "kpi_outcome": docx_frame_rows(kpi_outcome),
        "cost_benefit": docx_frame_rows(cost_benefit),
        "risks": docx_frame_rows(risks),
    }
    if st.button("Buat draft TOR dari template", type="primary"):
        document = make_tor_docx(
            selected_code,
            selected_row,
            hierarchy,
            adjustments,
        )
        st.session_state["tor_result"] = {
            "ro_code": selected_code,
            "activity": selected_row["Uraian"],
            "document": document,
        }

    result = st.session_state.get("tor_result")
    if result and result.get("ro_code") == selected_code:
        st.success("Draft TOR untuk RO {} siap diunduh.".format(selected_code))
        safe_name = re.sub(r"[^a-z0-9]+", "-", selected_row["Uraian"].lower()).strip("-") or "draft"
        safe_name += "-" + re.sub(r"[^a-z0-9]+", "-", selected_code.lower()).strip("-")
        st.download_button(
            "Unduh TOR dalam Word",
            data=result["document"],
            file_name="TOR-{}.docx".format(safe_name),
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


def main():
    st.markdown(
        """
        <style>
        .block-container { padding-top: 1.4rem; padding-bottom: 2rem; }
        h1, h2, h3 { color: #092759; }
        [data-testid="stMetric"] { background: #fff; border: 1px solid #e4ebf2;
            border-radius: 10px; padding: 14px 16px; }
        .stTabs [data-baseweb="tab-list"] { gap: 8px; }
        </style>
        """,
        unsafe_allow_html=True,
    )
    with st.sidebar:
        st.markdown("## Ruang Anggaran")
        logo_columns = st.columns(3)
        for column, filename in zip(logo_columns, ("balmon.png", "djid.png", "komdigi.png")):
            path = ROOT / "image" / filename
            if path.exists():
                column.image(str(path), use_container_width=True)
        st.caption("Balai Monitor SFR Kelas I Surabaya")
        st.divider()
        st.caption("Draft dan insight perlu ditinjau sebelum digunakan resmi.")

    st.title("Dashboard RAB & Generator TOR")
    data = load_analysis()
    dashboard_tab, tor_tab = st.tabs(["Dashboard", "Generator TOR"])
    with dashboard_tab:
        render_dashboard(data)
    with tor_tab:
        render_tor(data)


if __name__ == "__main__":
    main()
